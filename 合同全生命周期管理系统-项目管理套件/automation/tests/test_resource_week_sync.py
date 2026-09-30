from __future__ import annotations

import tempfile
import unittest
from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook

from ai_pmo.resource_week_sync import apply_resource_sync, plan_resource_sync


def save_plan_book(path: Path, missing: bool = False) -> None:
    book = Workbook()
    control = book.active
    control.title = "项目总控台账"
    control.append(["标题"])
    control.append(["说明"])
    control.append([f"列{i}" for i in range(1, 32)] + ["迭代周"])
    tasks = [
        ("T001", "汇编可研报告", "进行中", date(2026, 2, 2), date(2026, 2, 8)),
        ("T002", "盘点既有系统", "已完成", date(2026, 2, 9), date(2026, 2, 15)),
        ("T003", "访谈业务部门", "未开始", date(2026, 3, 1), date(2026, 3, 7)),
    ]
    for task_id, name, status, start, finish in tasks:
        row = [None] * 32
        row[0], row[4], row[6], row[8], row[9] = task_id, name, status, start, finish
        control.append(row)

    resource = book.create_sheet("资源计划")
    resource.append(["资源计划"])
    resource.append(["说明"])
    resource.append([])
    resource.append(["人员", "角色", "任务数", "未完成任务", "最早开始", "最晚结束",
                     "周可用工时", "计划投入工时", "负荷率", "负荷判断", "资源建议", "备注"])
    resource.append(["人员A-01", "人员A", 1, 1, datetime(2026, 2, 2), datetime(2026, 2, 8), 40, 40, None, None, None, None])
    resource.append(["人员B-01", "人员B", 1, 1, datetime(2026, 3, 1), datetime(2026, 3, 7), 40, 40, None, None, None, None])
    resource.append(["人员C-01", "人员C", 1, 1, datetime(2026, 3, 1), datetime(2026, 3, 7), 40, 40, None, None, None, None])
    for _ in range(resource.max_row, 37):
        resource.append([None])
    resource["A36"] = "资源周计划"
    resource.append(["人员", "周开始", "周结束", "是否有任务", "任务数", "周计划工时",
                     "负荷率", "负荷判断", "任务清单", None, None, "说明"])
    week_rows = [
        ("人员A-01", datetime(2026, 2, 2), datetime(2026, 2, 8), "有", 1, 40, "T001 旧任务名"),
        (None, datetime(2026, 2, 9), datetime(2026, 2, 15), "有", 2, 40, "T001 旧任务名；T002 旧名二"),
        ("人员B-01", datetime(2026, 3, 1), datetime(2026, 3, 7), "无", 0, 0, "无"),
        (None, datetime(2026, 3, 8), datetime(2026, 3, 14), "有", 1, 40, "T003 旧名三"),
        ("人员C-01", datetime(2026, 3, 1), datetime(2026, 3, 7), "有", 1, 40,
         "T001 旧任务名；T099 不存在" if missing else "T001 旧任务名"),
    ]
    for person, start, finish, has, count, hours, task_text in week_rows:
        resource.append([person, start, finish, has, count, hours, None, None, task_text, None, None, None])
    book.save(path)


class ResourceSyncPlanTests(unittest.TestCase):
    def test_plan_counts_rows_without_errors(self):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "plan.xlsx"
            save_plan_book(path)
            report = plan_resource_sync(path)
            self.assertEqual(report.errors, [])
            self.assertEqual(report.rows_refreshed, 4)
            self.assertEqual(len(report.persons), 3)

    def test_missing_task_reference_is_an_error(self):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "plan.xlsx"
            save_plan_book(path, missing=True)
            report = plan_resource_sync(path)
            self.assertIn("T099", report.errors[0])


class ResourceSyncApplyTests(unittest.TestCase):
    def test_apply_refreshes_task_text_counts_and_master(self):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "plan.xlsx"
            save_plan_book(path)
            report = apply_resource_sync(path)
            self.assertEqual(report.errors, [])

            book = load_workbook(path)
            resource = book["资源计划"]
            # 任务清单按总控当前名称刷新，多任务用中文分号
            self.assertEqual(resource["I39"].value, "T001 汇编可研报告")
            self.assertEqual(resource["I40"].value, "T001 汇编可研报告；T002 盘点既有系统")
            self.assertEqual(resource["E40"].value, 2)
            self.assertEqual(resource["F40"].value, 40)
            self.assertEqual(resource["I41"].value, "无")
            # master 汇总按周行重算：人员A 有任务2周→H=80
            self.assertEqual(resource["C5"].value, 2)
            self.assertEqual(resource["D5"].value, 1)
            self.assertEqual(resource["E5"].value, datetime(2026, 2, 2))
            self.assertEqual(resource["F5"].value, datetime(2026, 2, 15))
            self.assertEqual(resource["H5"].value, 80)
            self.assertEqual(resource["H6"].value, 40)

    def test_apply_refuses_when_task_missing(self):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "plan.xlsx"
            save_plan_book(path, missing=True)
            report = apply_resource_sync(path)
            self.assertTrue(report.errors)
            book = load_workbook(path)
            self.assertEqual(book["资源计划"]["I39"].value, "T001 旧任务名")

    def test_name_change_in_control_propagates(self):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "plan.xlsx"
            save_plan_book(path)
            apply_resource_sync(path)
            book = load_workbook(path)
            control = book["项目总控台账"]
            control.cell(row=4, column=5, value = "新任务名")
            book.save(path)
            apply_resource_sync(path)
            book = load_workbook(path)
            self.assertEqual(book["资源计划"]["I39"].value, "T001 新任务名")


if __name__ == "__main__":
    unittest.main()
