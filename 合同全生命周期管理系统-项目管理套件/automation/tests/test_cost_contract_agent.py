from __future__ import annotations

import tempfile
import unittest
from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook

from ai_pmo.agents.cost_contract import analyze_cost_contract


def save_cost_book(path: Path) -> None:
    book = Workbook()
    overview = book.active
    overview.title = "成本总览"
    overview.append([])

    baseline = book.create_sheet("预算基线")
    for _ in range(4):
        baseline.append([])
    baseline.append([
        "BUD-001", "建设成本", "项目人工", "一次性", "人天", 100, 1200,
        120000, 0, 120000, "项目经理", "已量化", "已批准",
    ])
    baseline.append([
        "BUD-002", "软件成本", "CA", "一次性", "许可", None, None,
        None, 0, None, "技术经理", "待估算", "等待POC",
    ])

    packages = book.create_sheet("采购包计划")
    for _ in range(4):
        packages.append([])
    packages.append([
        "PKG-001", "CA产品", "软件/服务", "许可与接口", "技术经理",
        datetime(2026, 9, 25), datetime(2026, 10, 15), None, "待估算", "POC后确认",
    ])

    book.create_sheet("合同台账")
    book.create_sheet("付款里程碑")
    book.create_sheet("成本偏差分析")
    book.save(path)


class CostContractAgentTests(unittest.TestCase):
    def test_flags_incomplete_budget_and_near_due_unpriced_package(self):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "cost.xlsx"
            save_cost_book(path)

            result = analyze_cost_contract(path, date(2026, 9, 18))

            self.assertEqual(
                {finding.object_id for finding in result.findings},
                {"BUDGET-COMPLETENESS", "PKG-001"},
            )
            budget = next(
                finding for finding in result.findings
                if finding.object_id == "BUDGET-COMPLETENESS"
            )
            self.assertEqual(budget.severity, "高")
            self.assertTrue(budget.requires_approval)


if __name__ == "__main__":
    unittest.main()
