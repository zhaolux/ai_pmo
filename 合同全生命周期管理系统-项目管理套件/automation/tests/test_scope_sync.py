from __future__ import annotations

import tempfile
import unittest
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook

from ai_pmo.scope_sync import (
    INTERFACE_EXECUTION_COLUMNS,
    INTERFACE_SCOPE_COLUMNS,
    SPEC_EXECUTION_COLUMNS,
    SPEC_SCOPE_COLUMNS,
    apply_sync,
    plan_sync,
)

INTERFACE_HEADER = list(INTERFACE_SCOPE_COLUMNS + INTERFACE_EXECUTION_COLUMNS)
SPEC_HEADER = list(SPEC_SCOPE_COLUMNS + SPEC_EXECUTION_COLUMNS)


def _pad(header: list[str], values: dict[str, object]) -> list[object]:
    return [values.get(name) for name in header]


def save_scope_book(path: Path, interface_rows: list[dict], spec_rows: list[dict]) -> None:
    book = Workbook()
    sheet = book.active
    sheet.title = "接口清单"
    for _ in range(3):
        sheet.append([])
    sheet.append(INTERFACE_HEADER)
    for values in interface_rows:
        sheet.append(_pad(INTERFACE_HEADER, values))

    specs = book.create_sheet("P0接口规格台账")
    for _ in range(3):
        specs.append([])
    specs.append(SPEC_HEADER)
    for values in spec_rows:
        specs.append(_pad(SPEC_HEADER, values))
    book.save(path)


def save_integration_book(path: Path, interface_rows: list[dict], spec_rows: list[dict]) -> None:
    book = Workbook()
    sheet = book.active
    sheet.title = "接口清单"
    for _ in range(3):
        sheet.append([])
    sheet.append(INTERFACE_HEADER)
    for values in interface_rows:
        sheet.append(_pad(INTERFACE_HEADER, values))

    specs = book.create_sheet("P0接口规格")
    for _ in range(3):
        specs.append([])
    specs.append(SPEC_HEADER)
    for values in spec_rows:
        specs.append(_pad(SPEC_HEADER, values))
    book.save(path)


def interface_row(key: str, name: str = "合同同步", priority: str = "P0",
                  spec_status: str = "待编制", test_status: str = "未开始") -> dict:
    return {
        "清单ID": key, "集成系统": "ERP", "接口名称": name, "接口描述": "同步合同",
        "接口方向": "出站", "范围决策": "纳入", "优先级": priority,
        "接口责任人": "集成工程师", "规格状态": spec_status,
        "测试状态": test_status, "备注": None,
    }


def spec_row(key: str, interface: str, status: str = "待编制") -> dict:
    values = {"规格ID": key, "清单ID": interface, "规格状态": status, "备注": None}
    for name in SPEC_SCOPE_COLUMNS:
        values.setdefault(name, f"{name}-{key}")
    return values


class ScopeSyncTests(unittest.TestCase):
    def _pair(self, folder: str):
        root = Path(folder)
        scope = root / "scope.xlsx"
        integration = root / "integration.xlsx"
        return scope, integration

    def test_identical_books_have_no_changes(self):
        with tempfile.TemporaryDirectory() as folder:
            scope, integration = self._pair(folder)
            save_scope_book(scope, [interface_row("IF-001")], [spec_row("SPEC-001", "IF-001")])
            save_integration_book(integration, [interface_row("IF-001")], [spec_row("SPEC-001", "IF-001")])

            report = plan_sync(scope, integration)

            self.assertFalse(any(p.has_changes for p in report.plans))
            self.assertFalse(report.plans[0].errors)

    def test_scope_drift_flows_to_integration_on_apply(self):
        with tempfile.TemporaryDirectory() as folder:
            scope, integration = self._pair(folder)
            save_scope_book(scope, [interface_row("IF-001", priority="P1")], [
                spec_row("SPEC-001", "IF-001"),
            ])
            save_integration_book(integration, [interface_row("IF-001", priority="P0")], [
                spec_row("SPEC-001", "IF-001"),
            ])

            report = apply_sync(scope, integration)

            self.assertTrue(report.applied)
            self.assertFalse(any(p.errors for p in report.plans))
            verify = plan_sync(scope, integration)
            self.assertFalse(any(p.has_changes for p in verify.plans))

    def test_execution_drift_flows_back_to_scope(self):
        with tempfile.TemporaryDirectory() as folder:
            scope, integration = self._pair(folder)
            save_scope_book(scope, [interface_row("IF-001", spec_status="待编制")], [
                spec_row("SPEC-001", "IF-001"),
            ])
            save_integration_book(
                integration,
                [interface_row("IF-001", spec_status="规格台账已建立", test_status="进行中")],
                [spec_row("SPEC-001", "IF-001", status="已评审")],
            )

            apply_sync(scope, integration)

            import openpyxl
            book = openpyxl.load_workbook(scope, data_only=True)
            row = [c.value for c in book["接口清单"][5]]
            values = dict(zip(INTERFACE_HEADER, row))
            self.assertEqual(values["规格状态"], "规格台账已建立")
            self.assertEqual(values["测试状态"], "进行中")
            spec_values = dict(zip(SPEC_HEADER, [c.value for c in book["P0接口规格台账"][5]]))
            self.assertEqual(spec_values["规格状态"], "已评审")

    def test_new_scope_record_is_appended_to_integration(self):
        with tempfile.TemporaryDirectory() as folder:
            scope, integration = self._pair(folder)
            save_scope_book(
                scope,
                [interface_row("IF-001"), interface_row("IF-002", name="付款同步")],
                [spec_row("SPEC-001", "IF-001"), spec_row("SPEC-002", "IF-002")],
            )
            save_integration_book(
                integration,
                [interface_row("IF-001")],
                [spec_row("SPEC-001", "IF-001")],
            )

            report = apply_sync(scope, integration)

            appended = report.plans[0]
            self.assertTrue(report.applied)
            verify = plan_sync(scope, integration)
            self.assertFalse(any(p.has_changes for p in verify.plans))

    def test_integration_only_record_is_error_not_deleted(self):
        with tempfile.TemporaryDirectory() as folder:
            scope, integration = self._pair(folder)
            save_scope_book(scope, [interface_row("IF-001")], [])
            save_integration_book(
                integration,
                [interface_row("IF-001"), interface_row("IF-099", name="幽灵接口")],
                [spec_row("SPEC-001", "IF-001"), spec_row("SPEC-099", "IF-099")],
            )

            report = apply_sync(scope, integration)

            self.assertTrue(any("IF-099" in e for p in report.plans for e in p.errors))
            self.assertTrue(any("SPEC-099" in e for p in report.plans for e in p.errors))

    def test_missing_column_is_error(self):
        with tempfile.TemporaryDirectory() as folder:
            scope, integration = self._pair(folder)
            save_scope_book(scope, [interface_row("IF-001")], [])
            book = Workbook()
            sheet = book.active
            sheet.title = "接口清单"
            for _ in range(3):
                sheet.append([])
            sheet.append(["清单ID", "优先级"])
            sheet.append(["IF-001", "P0"])
            book.create_sheet("P0接口规格")
            book.save(integration)

            report = plan_sync(scope, integration)

            self.assertTrue(any(p.errors for p in report.plans))


if __name__ == "__main__":
    unittest.main()
