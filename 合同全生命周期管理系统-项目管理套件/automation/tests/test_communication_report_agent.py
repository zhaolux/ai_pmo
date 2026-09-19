from __future__ import annotations

import tempfile
import unittest
from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook

from ai_pmo.agents.communication_report import analyze_communication_report


def save_communication_book(path: Path) -> None:
    book = Workbook()
    actions = book.active
    actions.title = "行动项台账"
    for _ in range(4):
        actions.append([])
    actions.append([
        "ACT-001", "会议", "MIN-001", "补充POC证据", "技术经理", None,
        datetime(2026, 9, 10), datetime(2026, 9, 17), "高", "进行中",
    ])

    decisions = book.create_sheet("决策日志")
    for _ in range(4):
        decisions.append([])
    decisions.append([
        "DEC-001", "CA产品POC选型", "技术/采购", datetime(2026, 9, 17),
        datetime(2026, 9, 30), "背景", "选项", "影响", "建议", "待决策",
        "项目治理委员会",
    ])
    book.create_sheet("升级事项")
    book.create_sheet("报告台账")
    book.save(path)


class CommunicationReportAgentTests(unittest.TestCase):
    def test_flags_overdue_action_and_due_decision(self):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "communication.xlsx"
            save_communication_book(path)

            result = analyze_communication_report(path, date(2026, 9, 18))

            self.assertEqual(
                {finding.object_id for finding in result.findings},
                {"ACT-001", "DEC-001"},
            )
            action = next(f for f in result.findings if f.object_id == "ACT-001")
            self.assertEqual(action.severity, "高")
            self.assertTrue(action.requires_approval)


if __name__ == "__main__":
    unittest.main()
