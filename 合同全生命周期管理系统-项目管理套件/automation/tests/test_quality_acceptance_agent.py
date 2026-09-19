from __future__ import annotations

import tempfile
import unittest
from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook

from ai_pmo.agents.quality_acceptance import analyze_quality_acceptance


def save_quality_book(path: Path) -> None:
    book = Workbook()
    overview = book.active
    overview.title = "质量总览"
    for _ in range(9):
        overview.append([])
    overview.append([
        "测试就绪门禁", "未通过", "测试负责人", datetime(2026, 9, 20),
        "环境、数据、用例就绪", "测试计划",
    ])

    defects = book.create_sheet("缺陷台账")
    for _ in range(4):
        defects.append([])
    defects.append([
        "BUG-001", datetime(2026, 9, 17), "系统测试", "REQ-001", "TS-001",
        "合同无法提交", "Critical", "主流程阻断", "高级开发人员",
        "代码", "修复", datetime(2026, 9, 19), "V1", "修复中",
    ])

    deliverables = book.create_sheet("交付物台账")
    for _ in range(4):
        deliverables.append([])
    deliverables.append([
        "D-001", "质量管理计划", "管理", "测试负责人",
        datetime(2026, 10, 31), None, "未开始", "项目经理", "待评审",
    ])
    book.create_sheet("UAT验收")
    book.save(path)


class QualityAcceptanceAgentTests(unittest.TestCase):
    def test_flags_due_gate_and_open_critical_defect(self):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "quality.xlsx"
            save_quality_book(path)

            result = analyze_quality_acceptance(path, date(2026, 9, 18))

            self.assertEqual(
                {finding.object_id for finding in result.findings},
                {"GATE-测试就绪门禁", "BUG-001"},
            )
            defect = next(f for f in result.findings if f.object_id == "BUG-001")
            self.assertEqual(defect.severity, "重大")
            self.assertTrue(defect.requires_approval)


if __name__ == "__main__":
    unittest.main()
