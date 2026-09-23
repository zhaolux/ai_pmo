import sys
import tempfile
import unittest
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook

script_dir = Path(__file__).parent.parent / 'scripts'
sys.path.insert(0, str(script_dir if script_dir.is_dir() else Path(__file__).parent))
from audit_migration_gates import audit_migration_gates, planned_sequence_issue


def book(path, sheets):
    wb = Workbook()
    wb.remove(wb.active)
    for name, rows in sheets.items():
        ws = wb.create_sheet(name)
        for row in rows:
            ws.append(row)
    wb.save(path)


class MigrationGateAuditTests(unittest.TestCase):
    def test_sequence_check_flags_reversed_stage_dates_without_calling_approval(self):
        self.assertIsNone(planned_sequence_issue('2027-05-14', ['2027-05-17', '2027-05-20'], '2027-05-29'))
        self.assertEqual(planned_sequence_issue('2027-05-20', ['2027-05-17'], '2027-05-29'),
                         'T087 演练完成日晚于 MIG 最早计划执行日，需核实阶段定义和依赖')

    def test_reads_authoritative_status_and_does_not_call_phases_conflicting(self):
        with tempfile.TemporaryDirectory() as temp:
            base = Path(temp)
            gate, plan, launch, migration, quality, risk = [base / f'{name}.xlsx' for name in ('gate', 'plan', 'launch', 'migration', 'quality', 'risk')]
            book(gate, {'启动门禁清单': [['G15', None, None, None, None, None, datetime(2026,10,13), '待确认']]})
            book(plan, {'项目总控台账': [
                ['T087', None, None, None, '迁移演练', None, '未开始', None, datetime(2027,5,13), datetime(2027,5,14)],
                ['T092', None, None, None, '生产切换', None, '未开始', None, datetime(2027,5,31), datetime(2027,5,31)],
            ]})
            book(launch, {
                '数据迁移与核验': [[f'MIG-{n:03d}', None, f'源{n}', None, None, None, datetime(2027,5,16+n), None, None, None, None, None, '未执行'] for n in range(1,7)],
                '上线准备清单': [['RDY-005', None, None, None, None, datetime(2027,5,20), None, None, None, None, '未就绪']],
                '切换执行计划': [['CUT-005', None, None, None, None, None, datetime(2027,5,29), datetime(2027,5,30), None, None, None, None, '未开始']],
            })
            book(migration, {'来源系统盘点': [['SRC-01', 'MIG-001', '现有合同系统/档案', '待盘点', '待确认', '待指定', None, None, None, None, None, '待盘点']]})
            book(quality, {'数据质量与迁移': [['数据质量与迁移验证']] + [[f'对象{n}', f'源{n}', None, None, '完整性', '准确性', '对账规则', datetime(2027,4,20), None, None, '未开始', None] for n in range(1,7)]})
            risk_row = [None] * 41
            risk_row[0], risk_row[1], risk_row[22], risk_row[27], risk_row[28], risk_row[40] = 'RSK-007', '历史合同数据质量不足', '产品经理', datetime(2026,10,13), '待处理', '数据责任矩阵与核对方案'
            book(risk, {'风险登记册': [risk_row]})
            result = audit_migration_gates(gate, plan, launch, migration, quality, risk)
            self.assertEqual(result['errors'], [])
            self.assertEqual(result['planned_sequence'], '演练 → 计划执行 → 生产切换')
            self.assertEqual(result['source_clue_count'], 1)
            self.assertIsNone(result['confirmed_system_count'])
            self.assertIn('G15 待确认', result['actions'])
            self.assertIn('来源系统待盘点', result['actions'])
            self.assertNotIn('日期冲突', ' '.join(result['actions']))
            self.assertEqual(result['quality']['planned_rule_coverage'], '6/6')
            self.assertEqual(result['quality']['verified_count'], 0)
            self.assertEqual(result['quality']['records'][0]['quality_source']['row'], 2)
            self.assertIn('迁移质量验证未完成 6 项', result['actions'])
            self.assertEqual(result['risk']['source']['row'], 1)
            self.assertEqual(result['risk']['owner'], '产品经理')
            self.assertEqual(result['risk']['status'], '待处理')
            self.assertIn('RSK-007 待处理', result['actions'])
            revised = load_workbook(risk)
            revised['风险登记册']['AC1'] = '已关闭'
            revised.save(risk)
            closed = audit_migration_gates(gate, plan, launch, migration, quality, risk)
            self.assertIn('RSK-007 关闭证据需人工核实', closed['actions'])

            partial = load_workbook(migration)
            partial['来源系统盘点']['D1'] = '合同数据库'
            partial['来源系统盘点']['L1'] = '已确认'
            partial.save(migration)
            incomplete = audit_migration_gates(gate, plan, launch, migration, quality, risk)
            self.assertIn('来源系统待盘点', incomplete['actions'])
            self.assertIn('数据责任人', incomplete['source_inventory']['records'][0]['missing_fields'])
            self.assertIn('合同/记录量', incomplete['source_inventory']['records'][0]['missing_fields'])

            filled = load_workbook(migration)
            row = filled['来源系统盘点']
            for column, value in {'E': '是', 'F': '负责人甲', 'G': 0, 'H': 0, 'I': '批量导出',
                                  'J': '保留查询', 'K': '盘点证据-01'}.items():
                row[f'{column}1'] = value
            filled.save(migration)
            fields_complete = audit_migration_gates(gate, plan, launch, migration, quality, risk)
            self.assertEqual(fields_complete['source_inventory']['records'][0]['missing_fields'], [])
            self.assertNotIn('来源系统待盘点', fields_complete['actions'])
            self.assertIsNone(fields_complete['confirmed_system_count'])

    def test_missing_required_id_is_structural_error(self):
        with tempfile.TemporaryDirectory() as temp:
            base = Path(temp)
            paths = [base / f'{name}.xlsx' for name in ('gate', 'plan', 'launch', 'migration')]
            book(paths[0], {'启动门禁清单': [['G15']]})
            book(paths[1], {'项目总控台账': [['T087']]})
            book(paths[2], {'数据迁移与核验': [['MIG-001']], '上线准备清单': [['RDY-005']], '切换执行计划': [['CUT-005']]})
            book(paths[3], {'来源系统盘点': [['SRC-01']]})
            result = audit_migration_gates(*paths)
            self.assertTrue(any('T092' in issue for issue in result['errors']))


if __name__ == '__main__':
    unittest.main()
