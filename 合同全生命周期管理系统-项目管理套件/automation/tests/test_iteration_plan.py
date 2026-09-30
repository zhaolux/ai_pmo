from __future__ import annotations

import tempfile
import unittest
from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook

from ai_pmo.iteration_plan import (
    DETAIL_START, SUMMARY_END, SUMMARY_START, apply_generation, plan_generation,
)


def save_plan_book(path: Path) -> None:
    book = Workbook()
    control = book.active
    control.title = "项目总控台账"
    headers = [f"列{i}" for i in range(1, 32)] + ["迭代周"]
    control.append(["标题"])
    control.append(["说明"])
    control.append(headers)
    rows = [
        ("T001", "任务一", "MS001", "已完成", "项目经理", None, date(2026, 9, 21), date(2026, 9, 25), "已完成", "2026-W39"),
        ("T002", "任务二", "MS001", "进行中", "产品经理", None, date(2026, 9, 22), date(2026, 9, 26), "进行中", "2026-W39"),
        ("T003", "任务三", "MS002", "未开始", "项目经理", None, date(2026, 9, 28), date(2026, 10, 2), "未开始", "2026-W40"),
        ("T004", "任务四", "MS002", "未开始", "技术经理", None, date(2026, 10, 5), date(2026, 10, 9), "未开始", "week-41"),
        ("T005", "任务五", "MS002", "未开始", "技术经理", None, date(2026, 10, 5), date(2026, 10, 9), "未开始", None),
    ]
    for item in rows:
        task_id, name, batch, status, owner, _pad, start, finish, biz_status, week = item
        row = [None] * 32
        row[0], row[4], row[5], row[6], row[7] = task_id, name, batch, status, owner
        row[8], row[9], row[12] = start, finish, biz_status
        row[31] = week
        control.append(row)

    iteration = book.create_sheet("周迭代计划")
    iteration.append(["周迭代计划与完成跟踪"])
    iteration.append(["说明"])
    iteration.append([])
    iteration.append(["迭代周", "开始日期", "结束日期", "周目标", "原周计划摘要", "计划完成阶段数",
                      "实际完成阶段数", "承诺完成率", "延期数", "暂停数", "周状态", "复盘结论"])
    iteration.append(["2026-W39", date(2026, 9, 21), date(2026, 9, 27), "本周目标", "=B31&\" \"&C31",
                      f"=COUNTIF($A${DETAIL_START}:$A$610,A5)", None, None, None, None,
                      f"=IF(J5>0,\"暂停\",IF(I5>0,\"延期\",IF(H5=1,\"已完成\",\"进行中\")))", "上周复盘"])
    while iteration.max_row < 30:
        iteration.append([None])
    iteration["A29"] = "周任务明细"
    iteration["A30"] = "迭代周"
    detail_row = DETAIL_START
    iteration.cell(row=detail_row, column=1, value="=$A$5")
    iteration.cell(row=detail_row, column=2, value="T000")
    iteration.cell(row=detail_row, column=3, value="=项目总控台账!E4")
    iteration.cell(row=detail_row, column=5, value="业务")
    iteration.cell(row=detail_row, column=11, value=date(2026, 9, 25))
    book.save(path)


class PlanGenerationTests(unittest.TestCase):
    def test_plan_groups_assigned_tasks_by_week(self):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "plan.xlsx"
            save_plan_book(path)
            report = plan_generation(path)
            self.assertEqual(report.errors, [])
            self.assertEqual(report.weeks, ["2026-W39", "2026-W40"])
            self.assertEqual(report.detail_count, 3)
            self.assertEqual(report.skipped, ["T004"])

    def test_week_capacity_overflow_is_an_error(self):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "plan.xlsx"
            save_plan_book(path)
            book = load_workbook(path)
            control = book["项目总控台账"]
            for index in range(30):
                row = [None] * 32
                row[0] = f"T{100 + index}"
                row[31] = f"2027-W{index + 1:02d}"
                control.append(row)
            book.save(path)
            report = plan_generation(path)
            self.assertTrue(any("超过" in error for error in report.errors))

    def test_plan_without_iteration_column_fails_clearly(self):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "plan.xlsx"
            save_plan_book(path)
            book = load_workbook(path)
            book["项目总控台账"].cell(row=3, column=32, value="别的列")
            book.save(path)
            report = plan_generation(path)
            self.assertTrue(any("迭代周" in error for error in report.errors))


class ApplyGenerationTests(unittest.TestCase):
    def test_apply_writes_formulas_and_preserves_manual_cells(self):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "plan.xlsx"
            save_plan_book(path)
            report = apply_generation(path)
            self.assertEqual(report.errors, [])

            book = load_workbook(path)
            iteration = book["周迭代计划"]
            # 明细区按周+任务ID排序生成，公式按任务ID INDEX/MATCH 查总控
            self.assertEqual(iteration.cell(row=DETAIL_START, column=1).value, "2026-W39")
            self.assertEqual(iteration.cell(row=DETAIL_START, column=2).value, "T001")
            formula = iteration.cell(row=DETAIL_START, column=3).value
            self.assertIn("INDEX", formula)
            self.assertIn("MATCH", formula)
            self.assertIn("$B31", formula)
            self.assertIn("项目总控台账!$A:$A", formula)
            self.assertNotIn("!E4", formula)
            # 周目标/复盘等人工列按周保留
            self.assertEqual(iteration["D5"].value, "本周目标")
            self.assertEqual(iteration["L5"].value, "上周复盘")
            # 新周行写入周一/周日日期
            self.assertEqual(iteration["A6"].value, "2026-W40")
            self.assertEqual(iteration["B6"].value, datetime(2026, 9, 28))
            self.assertEqual(iteration["C6"].value, datetime(2026, 10, 4))
            # 旧示例行 T000 被清除
            task_ids = [iteration.cell(row=r, column=2).value for r in range(DETAIL_START, DETAIL_START + 6)]
            self.assertNotIn("T000", task_ids)
            # 明细行数之外被清空
            self.assertIsNone(iteration.cell(row=DETAIL_START + 4, column=1).value)
            self.assertIsNone(iteration.cell(row=DETAIL_START + 4, column=2).value)

    def test_detail_preserves_stage_and_actual_date_by_task(self):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "plan.xlsx"
            save_plan_book(path)
            apply_generation(path)
            book = load_workbook(path)
            iteration = book["周迭代计划"]
            # 重新生成同一任务时保留其专业阶段与实际完成日期
            row_t1 = DETAIL_START
            iteration.cell(row=row_t1, column=5, value="后端")
            iteration.cell(row=row_t1, column=11, value=datetime(2026, 9, 26))
            book.save(path)
            apply_generation(path)
            book = load_workbook(path)
            iteration = book["周迭代计划"]
            self.assertEqual(iteration.cell(row=row_t1, column=5).value, "后端")
            self.assertEqual(iteration.cell(row=row_t1, column=11).value, datetime(2026, 9, 26))

    def test_unassigned_task_row_is_cleared_on_regeneration(self):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "plan.xlsx"
            save_plan_book(path)
            apply_generation(path)
            book = load_workbook(path)
            control = book["项目总控台账"]
            for row in range(4, 10):
                if control.cell(row=row, column=1).value == "T002":
                    control.cell(row=row, column=32).value = None
            book.save(path)
            apply_generation(path)
            book = load_workbook(path)
            iteration = book["周迭代计划"]
            task_ids = [iteration.cell(row=r, column=2).value for r in range(DETAIL_START, DETAIL_START + 3)]
            self.assertEqual(task_ids, ["T001", "T003", None])
            self.assertIsNone(iteration.cell(row=DETAIL_START + 1, column=5).value)

    def test_apply_refuses_on_capacity_error(self):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "plan.xlsx"
            save_plan_book(path)
            book = load_workbook(path)
            control = book["项目总控台账"]
            for index in range(30):
                row = [None] * 32
                row[0] = f"T{100 + index}"
                row[31] = f"2027-W{index + 1:02d}"
                control.append(row)
            book.save(path)
            report = apply_generation(path)
            self.assertTrue(report.errors)
            book = load_workbook(path)
            self.assertEqual(book["周迭代计划"].cell(row=DETAIL_START, column=2).value, "T000")


if __name__ == "__main__":
    unittest.main()
