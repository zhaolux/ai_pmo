from __future__ import annotations

import json
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook

from scripts import sync_cost_caliber_to_g15 as g15_sync


def make_suite(root: Path) -> Path:
    cost_dir = root / "05_成本与合同管理"
    cost_dir.mkdir(parents=True)
    cost = cost_dir / "成本与合同管理-20260918-V5.xlsx"
    book = Workbook()
    sheet = book.active
    sheet.title = "人工成本预算"
    sheet["A15"] = "合计"
    sheet["D15"] = 4721
    sheet["F15"] = 6525700
    budget = book.create_sheet("预算基线")
    budget["J5"] = 6525700
    budget["J6"] = 9000000
    book.save(cost)
    registry = root / "automation" / "config" / "current_workbooks.json"
    registry.parent.mkdir(parents=True)
    registry.write_text(json.dumps({
        "cost": cost.relative_to(root).as_posix(),
    }, ensure_ascii=False), encoding="utf-8")
    return cost


def make_g15_xlsx(path: Path, text: str) -> None:
    book = Workbook()
    sheet = book.active
    sheet.title = "签认跟踪"
    sheet["H8"] = text
    book.save(path)


class CaliberSyncPreviewTests(unittest.TestCase):
    def test_preview_reports_targets_without_writing(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            make_suite(root)
            g15 = root / "01_项目启动与治理"
            g15.mkdir(parents=True)
            make_g15_xlsx(
                g15 / "G15-会前签认跟踪-20260924-V1.xlsx",
                "确认计划人天唯一来源为《月度投入明细》：5,433 人天 / 769.77 万元 / 均价 1,416.84 元（G15 会前完成）",
            )
            (g15 / "G15-未决事项处置建议包-20260923.md").write_text(
                "现状：**5,433 人天 / 769.77 万元 / 平均人日单价 1,416.84 元**\n"
                "预算基线以“人工成本 769.77 万（月度投入明细驱动）”\n"
                "证据：5,433 人天/769.77 万元/平均人日单价 1,416.84 元\n"
                "确认\"5,433 人天/769.77 万元（月度投入明细驱动）\"为批准基线\n",
                encoding="utf-8",
            )
            (g15 / "G15-会前材料包清单与缺口对照-20260923.md").write_text(
                "驾驶舱 2026-09-23 口径为 5,433 人天/769.77 万元/均价 1,416.84 元\n",
                encoding="utf-8",
            )

            preview = g15_sync.preview_sync(root)

            self.assertEqual(preview["metrics"].person_days, 4721)
            plan = preview["plan"]
            self.assertEqual(len(plan), 3)
            self.assertTrue(all(item["replacements"] > 0 for item in plan))
            # 预览不写盘
            xlsx = g15 / "G15-会前签认跟踪-20260924-V1.xlsx"
            from openpyxl import load_workbook
            self.assertIn("5,433", load_workbook(xlsx)["签认跟踪"]["H8"].value)
            self.assertFalse((g15 / "_archive").exists())

    def test_preview_with_no_g15_files_reports_empty_plan(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            make_suite(root)
            preview = g15_sync.preview_sync(root)
            self.assertEqual(preview["plan"], [])

    def test_preview_raises_when_cost_total_row_missing(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            cost = make_suite(root)
            from openpyxl import load_workbook
            book = load_workbook(cost)
            book["人工成本预算"]["A15"] = "别的"
            book.save(cost)
            with self.assertRaises(ValueError):
                g15_sync.preview_sync(root)


class CaliberSyncApplyTests(unittest.TestCase):
    def test_apply_updates_xlsx_and_md_with_backup(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            make_suite(root)
            g15 = root / "01_项目启动与治理"
            g15.mkdir(parents=True)
            xlsx = g15 / "G15-会前签认跟踪-20260924-V1.xlsx"
            make_g15_xlsx(
                xlsx,
                "确认计划人天唯一来源为《月度投入明细》：5,433 人天 / 769.77 万元 / 均价 1,416.84 元（G15 会前完成）",
            )
            md = g15 / "G15-未决事项处置建议包-20260923.md"
            md.write_text(
                "现状：**5,433 人天 / 769.77 万元 / 平均人日单价 1,416.84 元**\n"
                "预算基线以“人工成本 769.77 万（月度投入明细驱动）”\n"
                "证据：5,433 人天/769.77 万元/平均人日单价 1,416.84 元\n"
                "确认\"5,433 人天/769.77 万元（月度投入明细驱动）\"为批准基线\n",
                encoding="utf-8",
            )

            g15_sync.sync(root)

            from openpyxl import load_workbook
            text = load_workbook(xlsx)["签认跟踪"]["H8"].value
            self.assertIn("4,721 人天 / 652.57 万元 / 均价 1,382.27 元", text)
            self.assertNotIn("5,433", text)
            updated = md.read_text(encoding="utf-8")
            self.assertIn("4,721 人天 / 652.57 万元 / 平均人日单价 1,382.27 元", updated)
            self.assertIn("人工成本 652.57 万（月度投入明细驱动）", updated)
            self.assertIn('"4,721 人天/652.57 万元（月度投入明细驱动）"', updated)
            self.assertNotIn("5,433", updated)
            # 备份存在且保留旧内容
            backups = list((g15 / "_archive").rglob("*签认跟踪*.xlsx"))
            self.assertEqual(len(backups), 1)
            old_text = load_workbook(backups[0])["签认跟踪"]["H8"].value
            self.assertIn("5,433", old_text)

    def test_apply_updates_dashboard_date_line(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            make_suite(root)
            g15 = root / "01_项目启动与治理"
            g15.mkdir(parents=True)
            md = g15 / "G15-会前材料包清单与缺口对照-20260923.md"
            md.write_text(
                "驾驶舱 2026-09-23 口径为 5,433 人天/769.77 万元/均价 1,416.84 元\n",
                encoding="utf-8",
            )
            g15_sync.sync(root)
            updated = md.read_text(encoding="utf-8")
            self.assertIn("驾驶舱 20", updated)
            self.assertIn("4,721 人天/652.57 万元/均价 1,382.27 元", updated)
            self.assertNotIn("5,433", updated)


if __name__ == "__main__":
    unittest.main()
