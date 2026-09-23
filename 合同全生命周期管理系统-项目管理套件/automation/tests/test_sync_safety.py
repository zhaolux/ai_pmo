from __future__ import annotations

import json
import tempfile
import unittest
from dataclasses import fields
from datetime import date
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook

from scripts import sync_all_to_ai_pmo as sync


def register(suite: Path, values: dict[str, Path]) -> None:
    registry = suite / "automation" / "config" / "current_workbooks.json"
    registry.parent.mkdir(parents=True, exist_ok=True)
    registry.write_text(json.dumps({
        key: path.relative_to(suite).as_posix() for key, path in values.items()
    }, ensure_ascii=False), encoding="utf-8")


class SyncSafetyTests(unittest.TestCase):
    def test_preview_lists_actual_target_and_sources_without_writing(self):
        with tempfile.TemporaryDirectory() as folder:
            suite = Path(folder)
            files = {
                "02_计划与进度管理": "计划与进度管理-V1.xlsx",
                "01_项目启动与治理": "项目启动与治理-V1.xlsx",
                "06_风险问题与变更": "风险管理-V1.xlsx",
                "11_AI_PMO中心": "AI PMO中心-V1.xlsx",
            }
            for directory, name in files.items():
                path = suite / directory / name
                path.parent.mkdir(parents=True)
                path.write_bytes(b"existing")
            (suite / "06_风险问题与变更" / "问题与变更管理-V1.xlsx").write_bytes(b"existing")
            other = suite / "11_AI_PMO中心" / "AI PMO-Agent运行结果-V1.xlsx"
            other.write_bytes(b"newer-but-not-target")
            register(suite, {
                "plan": suite / "02_计划与进度管理" / files["02_计划与进度管理"],
                "gate": suite / "01_项目启动与治理" / files["01_项目启动与治理"],
                "risk": suite / "06_风险问题与变更" / files["06_风险问题与变更"],
                "change": suite / "06_风险问题与变更" / "问题与变更管理-V1.xlsx",
                "ai_pmo": suite / "11_AI_PMO中心" / files["11_AI_PMO中心"],
            })

            preview = sync.preview_sync(suite, date(2026, 9, 20))

            self.assertEqual(preview["target"], suite / "11_AI_PMO中心" / files["11_AI_PMO中心"])
            self.assertEqual(preview["sources"]["plan"].name, files["02_计划与进度管理"])
            self.assertEqual(preview["sources"]["risk"].name, files["06_风险问题与变更"])
            self.assertEqual(preview["data_date"], "2026-09-20")
            self.assertEqual(preview["target"].read_bytes(), b"existing")
            self.assertFalse((suite / "11_AI_PMO中心" / "_archive").exists())

    def test_preview_ignores_newer_unregistered_ai_pmo_version(self):
        with tempfile.TemporaryDirectory() as folder:
            suite = Path(folder)
            target = suite / "11_AI_PMO中心" / "AI PMO中心-20260920-V1.xlsx"
            target.parent.mkdir(parents=True)
            target.write_bytes(b"registered")
            newer = target.parent / "AI PMO中心-20260921-V1.xlsx"
            newer.write_bytes(b"unapproved")
            values = {}
            for key, directory, name in (
                ("plan", "02_计划与进度管理", "计划与进度管理"),
                ("gate", "01_项目启动与治理", "项目启动与治理"),
                ("risk", "06_风险问题与变更", "风险管理"),
                ("change", "06_风险问题与变更", "问题与变更管理"),
            ):
                path = suite / directory / f"{name}-20260920-V1.xlsx"
                path.parent.mkdir(parents=True, exist_ok=True)
                path.write_bytes(b"registered")
                values[key] = path
            values["ai_pmo"] = target
            register(suite, values)

            self.assertEqual(sync.preview_sync(suite, date(2026, 9, 20))["target"], target)

    def test_backup_is_byte_identical_and_never_reuses_path(self):
        with tempfile.TemporaryDirectory() as folder:
            target = Path(folder) / "11_AI_PMO中心" / "AI PMO中心.xlsx"
            target.parent.mkdir(parents=True)
            target.write_bytes(b"human-confirmed")

            first = sync.backup_workbook(target)
            second = sync.backup_workbook(target)

            self.assertNotEqual(first, second)
            self.assertEqual(first.read_bytes(), b"human-confirmed")
            self.assertEqual(second.read_bytes(), b"human-confirmed")
            self.assertEqual(target.read_bytes(), b"human-confirmed")

    def test_log_row_skips_partially_filled_human_row(self):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "AI PMO.xlsx"
            book = Workbook()
            sheet = book.active
            sheet.title = "AI运行日志"
            sheet["A5"] = "人工备注"
            sheet["C6"] = "=1+1"
            book.save(path)

            row, ids = sync.find_next_log_row(path)

            self.assertEqual(row, 7)
            self.assertIn("人工备注", ids)

    def test_watchlist_manual_columns_require_same_task_id_at_same_row(self):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "AI PMO.xlsx"
            book = Workbook()
            sheet = book.active
            sheet.title = "两周督办"
            sheet["B5"] = "T001"
            sheet["L5"] = "2026-09-25"
            sheet["M5"] = "已提交验收记录"
            book.save(path)
            original = path.read_bytes()

            sync.ensure_watchlist_manual_alignment(path, [{"task_id": "T001"}])
            with self.assertRaisesRegex(RuntimeError, "T001"):
                sync.ensure_watchlist_manual_alignment(path, [{"task_id": "T002"}])
            with self.assertRaisesRegex(RuntimeError, "T001"):
                sync.ensure_watchlist_manual_alignment(path, [])
            self.assertEqual(path.read_bytes(), original)

    def test_watchlist_reorder_without_manual_columns_is_allowed(self):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "AI PMO.xlsx"
            book = Workbook()
            sheet = book.active
            sheet.title = "两周督办"
            sheet["B5"] = "T001"
            book.save(path)

            sync.ensure_watchlist_manual_alignment(path, [{"task_id": "T002"}])

    def test_shorter_watchlist_identifies_stale_generated_rows(self):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "AI PMO.xlsx"
            book = Workbook()
            sheet = book.active
            sheet.title = "两周督办"
            sheet["B5"] = "T001"
            sheet["B6"] = "T002"
            sheet["C8"] = "无编号旧内容"
            book.save(path)

            self.assertEqual(sync.find_stale_watchlist_rows(path, new_count=1), [6, 8])

    def test_main_previews_by_default_without_calling_sync(self):
        preview = {
            "target": Path("/tmp/target.xlsx"),
            "sources": {"plan": Path("/tmp/plan.xlsx")},
            "data_date": "2026-09-20",
        }
        with patch("sys.argv", ["sync_all_to_ai_pmo.py", "--suite", "/tmp/suite"]), patch.object(
            sync, "preview_sync", return_value=preview
        ), patch.object(sync, "sync_ai_pmo") as execute:
            sync.main()
            execute.assert_not_called()

    def test_apply_stops_if_target_changes_after_preview(self):
        with tempfile.TemporaryDirectory() as folder:
            suite = Path(folder)
            directory = suite / "11_AI_PMO中心"
            directory.mkdir()
            expected = directory / "AI PMO中心-旧版.xlsx"
            expected.write_bytes(b"old")
            current = directory / "AI PMO中心-新版.xlsx"
            current.write_bytes(b"new")
            register(suite, {"ai_pmo": current})

            with self.assertRaisesRegex(RuntimeError, "预览后已变化"):
                sync.sync_ai_pmo(suite, date(2026, 9, 20), expected_target=expected)

            self.assertEqual(expected.read_bytes(), b"old")
            self.assertEqual(current.read_bytes(), b"new")

    def test_watchlist_ignores_newer_unrelated_workbook(self):
        with tempfile.TemporaryDirectory() as folder:
            suite = Path(folder)
            directory = suite / "02_计划与进度管理"
            directory.mkdir()
            plan = Workbook()
            plan.active.title = "项目总控台账"
            plan.save(directory / "计划与进度管理-V1.xlsx")
            unrelated = Workbook()
            unrelated.active.title = "其他"
            unrelated.save(directory / "临时分析.xlsx")
            register(suite, {"plan": directory / "计划与进度管理-V1.xlsx"})

            self.assertEqual(sync.extract_two_week_watchlist(suite, date(2026, 9, 20)), [])

    def test_quality_check_ignores_agent_result_as_module_workbook(self):
        with tempfile.TemporaryDirectory() as folder:
            suite = Path(folder)
            plan_dir = suite / "02_计划与进度管理"
            plan_dir.mkdir()
            plan = Workbook()
            plan.active.title = "项目总控台账"
            plan.save(plan_dir / "计划与进度管理-V1.xlsx")
            center_dir = suite / "11_AI_PMO中心"
            center_dir.mkdir()
            center = Workbook()
            center.active.title = "_数据接口"
            center.active["B4"] = 11
            center.save(center_dir / "AI PMO中心-V1.xlsx")
            agent = Workbook()
            agent.active.title = "Agent运行结果"
            agent.save(center_dir / "AI PMO-Agent运行结果-V1.xlsx")
            register(suite, {
                "plan": plan_dir / "计划与进度管理-V1.xlsx",
                "ai_pmo": center_dir / "AI PMO中心-V1.xlsx",
            })

            issues = sync.collect_data_quality_issues(suite, date(2026, 9, 20))

            self.assertFalse(any(
                issue["source"] == "AI PMO-Agent运行结果-V1.xlsx"
                for issue in issues
            ))

    def test_run_log_advice_does_not_claim_blockers_when_none_exist(self):
        metrics = sync.ProjectMetrics(**{field.name: 0 for field in fields(sync.ProjectMetrics)})
        issues = [{"status": "需整改"}]

        record = sync.build_run_log_record("RUN-1", date(2026, 9, 20), metrics, issues)

        self.assertNotIn("阻断项", record["advice"])
        self.assertIn("需整改", record["advice"])

    def test_main_apply_uses_previewed_target(self):
        target = Path("/tmp/target.xlsx")
        preview = {
            "target": target,
            "sources": {"plan": Path("/tmp/plan.xlsx")},
            "data_date": "2026-09-20",
        }
        metrics = type("Metrics", (), {
            "total_tasks": 1, "overdue_tasks": 0,
            "major_residual_risks": 0, "high_residual_risks": 0,
        })()
        with patch("sys.argv", ["sync_all_to_ai_pmo.py", "--suite", "/tmp/suite", "--date", "2026-09-20", "--apply"]), patch.object(
            sync, "preview_sync", return_value=preview
        ), patch.object(sync, "sync_ai_pmo", return_value=(target, metrics, Path("/tmp/backup.xlsx"))) as execute:
            sync.main()
            execute.assert_called_once_with(
                Path("/tmp/suite").resolve(), date(2026, 9, 20), expected_target=target,
            )

    def test_refresh_command_guard_rejects_manual_diagnosis_cells(self):
        with self.assertRaisesRegex(ValueError, "人工维护区域"):
            sync.validate_refresh_commands([
                {"command": "set", "path": "/AI项目诊断/B11", "props": {"value": "绿"}},
            ], log_row=5)

        sync.validate_refresh_commands([
            {"command": "set", "path": "/AI项目诊断/C11", "props": {"formula": "1"}},
            {"command": "set", "path": "/AI运行日志/I5", "props": {"value": "待确认"}},
        ], log_row=5)

    def test_full_sync_generated_commands_pass_manual_region_guard(self):
        with tempfile.TemporaryDirectory() as folder:
            suite = Path(folder)
            names = {
                "02_计划与进度管理": "计划与进度管理-V1.xlsx",
                "01_项目启动与治理": "项目启动与治理-V1.xlsx",
                "06_风险问题与变更": "风险管理-V1.xlsx",
                "11_AI_PMO中心": "AI PMO中心-V1.xlsx",
            }
            for directory, name in names.items():
                path = suite / directory / name
                path.parent.mkdir(parents=True)
                path.write_bytes(b"fixture")
            target = suite / "11_AI_PMO中心" / names["11_AI_PMO中心"]
            register(suite, {
                "plan": suite / "02_计划与进度管理" / names["02_计划与进度管理"],
                "gate": suite / "01_项目启动与治理" / names["01_项目启动与治理"],
                "risk": suite / "06_风险问题与变更" / names["06_风险问题与变更"],
                "ai_pmo": target,
            })
            metrics = sync.ProjectMetrics(**{field.name: 0 for field in fields(sync.ProjectMetrics)})
            with patch.object(sync, "extract_project_metrics", return_value=metrics), patch.object(
                sync, "extract_two_week_watchlist", return_value=[]
            ), patch.object(sync, "extract_milestone_overview", return_value=[]
            ), patch.object(sync, "extract_management_focus", return_value=[]
            ), patch.object(sync, "collect_data_quality_issues", return_value=[]), patch.object(
                sync, "find_next_log_row", return_value=(5, [])
            ), patch.object(sync, "ensure_watchlist_manual_alignment"
            ), patch.object(sync, "find_stale_watchlist_rows", return_value=[6]
            ), patch.object(sync, "backup_workbook", return_value=Path(folder) / "backup.xlsx"), patch.object(
                sync.shutil, "which", return_value="officecli"
            ), patch.object(sync.subprocess, "run", side_effect=lambda args, **kwargs: commands.extend(
                json.loads(Path(args[args.index("--input") + 1]).read_text(encoding="utf-8"))
            )) as run:
                commands = []
                result_target, _, _ = sync.sync_ai_pmo(suite, date(2026, 9, 20), expected_target=target)

            self.assertEqual(result_target, target)
            run.assert_called_once()
            self.assertEqual(
                {command["path"] for command in commands if command["path"].startswith("/两周督办/")},
                {f"/两周督办/{column}6" for column in "ABCDEFGHIJK"},
            )

    def test_manual_watchlist_conflict_stops_before_backup(self):
        with tempfile.TemporaryDirectory() as folder:
            suite = Path(folder)
            target = suite / "11_AI_PMO中心" / "AI PMO中心-V1.xlsx"
            target.parent.mkdir(parents=True)
            target.write_bytes(b"fixture")
            register(suite, {"ai_pmo": target})
            metrics = sync.ProjectMetrics(**{field.name: 0 for field in fields(sync.ProjectMetrics)})
            with patch.object(sync, "extract_project_metrics", return_value=metrics), patch.object(
                sync, "extract_two_week_watchlist", return_value=[{"task_id": "T002"}]
            ), patch.object(sync, "extract_milestone_overview", return_value=[]
            ), patch.object(sync, "extract_management_focus", return_value=[]
            ), patch.object(sync, "collect_data_quality_issues", return_value=[]), patch.object(
                sync, "find_next_log_row", return_value=(5, [])
            ), patch.object(sync, "ensure_watchlist_manual_alignment", side_effect=RuntimeError("人工列错配")), patch.object(
                sync, "backup_workbook"
            ) as backup, patch.object(sync.subprocess, "run") as run:
                with self.assertRaisesRegex(RuntimeError, "人工列错配"):
                    sync.sync_ai_pmo(suite, date(2026, 9, 20), expected_target=target)
            backup.assert_not_called()
            run.assert_not_called()

    def _make_migration_workbook(self, suite: Path) -> Path:
        path = suite / "09_上线移交与运营" / "旧合同系统迁移管理-V1.xlsx"
        path.parent.mkdir(parents=True, exist_ok=True)
        book = Workbook()
        sheet = book.active
        sheet.title = "_数据接口"
        for row_number, (label, value) in enumerate([
            ("字段", "值"), ("Project_ID", "ECM-2026"), ("模块编码", "09"),
            ("来源线索数", 6), ("已盘点来源系统数", 2), ("迁移门禁数", 7),
            ("G15 目标日期", date(2026, 10, 13)),
        ], start=1):
            sheet.cell(row_number, 1, label)
            sheet.cell(row_number, 2, value)
        book.save(path)
        return path

    def test_extract_migration_metrics_reads_registered_interface(self):
        with tempfile.TemporaryDirectory() as folder:
            suite = Path(folder)
            migration = self._make_migration_workbook(suite)
            register(suite, {"migration": migration})
            metrics = sync.extract_migration_metrics(suite)
            self.assertEqual(metrics, {
                "来源线索数": 6, "已盘点来源系统数": 2,
                "迁移门禁数": 7, "G15目标日期": "2026-10-13",
            })

    def test_extract_migration_metrics_is_none_without_registration(self):
        with tempfile.TemporaryDirectory() as folder:
            self.assertIsNone(sync.extract_migration_metrics(Path(folder)))

    def test_extract_milestone_overview_reads_first_five_milestones(self):
        with tempfile.TemporaryDirectory() as folder:
            suite = Path(folder)
            path = suite / "02_计划与进度管理" / "计划与进度管理-V1.xlsx"
            path.parent.mkdir(parents=True)
            book = Workbook()
            sheet = book.active
            sheet.title = "里程碑计划"
            sheet.append(["标题"] * 14)
            sheet.append(["表头"] * 14)
            sheet.append(["说明"] * 14)
            for index in range(7):
                sheet.append([f"MS00{index+1}", "", f"里程碑{index+1}", "", f"负责人{index+1}",
                              "", "", date(2026, 2, index + 1), "", "", "进行中", "", "绿", "交付物"])
            book.save(path)
            register(suite, {"plan": path})
            overview = sync.extract_milestone_overview(suite)
            self.assertEqual(len(overview), 5)
            self.assertEqual(overview[0], {
                "name": "里程碑1", "due": date(2026, 2, 1),
                "status": "进行中", "health": "绿", "owner": "负责人1",
            })
            self.assertEqual(overview[-1]["name"], "里程碑5")

    def test_extract_management_focus_major_first_then_due_date(self):
        with tempfile.TemporaryDirectory() as folder:
            suite = Path(folder)
            risk = suite / "06_风险问题与变更" / "风险管理-V1.xlsx"
            risk.parent.mkdir(parents=True)
            book = Workbook()
            sheet = book.active
            sheet.title = "风险登记册"
            for _ in range(8):
                sheet.append(["表头"] * 44)
            rows = [
                ("RSK-001", "高风险甲", "减轻", date(2026, 9, 25), "高", None),
                ("RSK-004", "重大风险乙", "规避", date(2026, 10, 13), "重大", None),
                ("RSK-010", "中风险丙", "减轻", date(2026, 9, 20), "中", None),
                ("RSK-011", "已关闭重大", "减轻", date(2026, 9, 1), "重大", date(2026, 9, 2)),
            ]
            for r in rows:
                row = [""] * 44
                row[0], row[1], row[24], row[27], row[33] = r[0], r[1], r[2], r[3], r[4]
                row[22], row[39] = "责任人" + r[0][-3:], r[5]
                sheet.append(row)
            book.save(risk)
            change = suite / "06_风险问题与变更" / "问题与变更管理-V1.xlsx"
            book2 = Workbook()
            sheet2 = book2.active
            sheet2.title = "待决策事项"
            for _ in range(4):
                sheet2.append(["表头"] * 13)
            sheet2.append(["D-001", "", "决策主题甲", "", "", "建议方案甲", "委员会甲",
                           date(2026, 9, 25), "待决策", "", "", "", ""])
            sheet2.append(["D-002", "", "决策主题乙", "", "", "建议方案乙", "委员会乙",
                           date(2026, 10, 9), "已决策", "", "", "", ""])
            book2.save(change)
            register(suite, {"risk": risk, "change": change})
            focus = sync.extract_management_focus(suite)
            self.assertEqual([item["item"] for item in focus],
                             ["决策主题甲", "重大风险乙", "高风险甲"])
            self.assertEqual(focus[0]["level"], "重大")
            self.assertIn("建议方案甲", focus[0]["next"])
            self.assertEqual(focus[1]["owner"], "责任人004")
            self.assertIn("最迟 2026-09-25", focus[2]["next"])


    def test_full_sync_writes_migration_interface_rows_and_dashboard_cells(self):
        with tempfile.TemporaryDirectory() as folder:
            suite = Path(folder)
            names = {
                "02_计划与进度管理": "计划与进度管理-V1.xlsx",
                "01_项目启动与治理": "项目启动与治理-V1.xlsx",
                "06_风险问题与变更": "风险管理-V1.xlsx",
                "11_AI_PMO中心": "AI PMO中心-V1.xlsx",
            }
            for directory, name in names.items():
                path = suite / directory / name
                path.parent.mkdir(parents=True)
                path.write_bytes(b"fixture")
            target = suite / "11_AI_PMO中心" / names["11_AI_PMO中心"]
            migration = self._make_migration_workbook(suite)
            register(suite, {
                "plan": suite / "02_计划与进度管理" / names["02_计划与进度管理"],
                "gate": suite / "01_项目启动与治理" / names["01_项目启动与治理"],
                "risk": suite / "06_风险问题与变更" / names["06_风险问题与变更"],
                "ai_pmo": target,
                "migration": migration,
            })
            metrics = sync.ProjectMetrics(**{field.name: 0 for field in fields(sync.ProjectMetrics)})
            milestones = [{"name": "可研工作启动", "due": date(2026, 2, 10), "status": "已完成",
                           "health": "绿", "owner": "项目经理"}]
            with patch.object(sync, "extract_project_metrics", return_value=metrics), patch.object(
                sync, "extract_two_week_watchlist", return_value=[]
            ), patch.object(sync, "extract_milestone_overview", return_value=milestones
            ), patch.object(sync, "extract_management_focus", return_value=[{"item": "POC选型", "next": "9月25日前完成", "owner": "技术经理", "level": "重大"}]
            ), patch.object(sync, "collect_data_quality_issues", return_value=[]), patch.object(
                sync, "find_next_log_row", return_value=(5, [])
            ), patch.object(sync, "ensure_watchlist_manual_alignment"
            ), patch.object(sync, "find_stale_watchlist_rows", return_value=[]), patch.object(
                sync, "backup_workbook", return_value=Path(folder) / "backup.xlsx"), patch.object(
                sync.shutil, "which", return_value="officecli"
            ), patch.object(sync.subprocess, "run", side_effect=lambda args, **kwargs: commands.extend(
                json.loads(Path(args[args.index("--input") + 1]).read_text(encoding="utf-8"))
            )):
                commands = []
                sync.sync_ai_pmo(suite, date(2026, 9, 20), expected_target=target)

            interface_paths = {command["path"] for command in commands if command["path"].startswith("/_数据接口/")}
            self.assertTrue({"/_数据接口/A39", "/_数据接口/B39", "/_数据接口/B40", "/_数据接口/B41", "/_数据接口/B42"} <= interface_paths)
            dashboard_paths = {command["path"] for command in commands if command["path"].startswith("/项目驾驶舱/")}
            self.assertTrue({"/项目驾驶舱/I18", "/项目驾驶舱/I19", "/项目驾驶舱/J19", "/项目驾驶舱/K19", "/项目驾驶舱/L19"} <= dashboard_paths)
            self.assertTrue({f"/项目驾驶舱/{column}{row}" for column in "IJKLM" for row in range(13, 17)} <= dashboard_paths)
            self.assertIn("/项目驾驶舱/I12", dashboard_paths)
            self.assertTrue({f"/项目驾驶舱/{column}{row}" for column in "ABCDE" for row in range(21, 25)} <= dashboard_paths)
            self.assertIn("/项目驾驶舱/B20", dashboard_paths)
            sync.validate_refresh_commands(commands, log_row=5)


if __name__ == "__main__":
    unittest.main()
