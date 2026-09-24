from __future__ import annotations

import json
import tempfile
import unittest
from pathlib import Path

from docx import Document
from openpyxl import Workbook

from ai_pmo.weekly_report import (
    build_weekly_docx,
    extract_cost_overview,
    weekly_report_payload,
)


def make_report() -> dict:
    return {
        "data_date": "2026-09-24",
        "health": "黄",
        "findings": [],
        "agent_results": [{"agent": "成本与合同", "finding_count": 2, "summary": "识别2项成本关注"}],
        "management_summary": {
            "finding_count": 5, "major_count": 1, "high_count": 2, "medium_count": 2,
            "approval_count": 1,
            "top_priorities": [
                {"severity": "高", "title": "D-001 产品选型", "object_id": "DEC-001",
                 "owner": "指导委员会", "recommendation": "按类别择优入围"}
            ],
            "decisions_needed": [
                {"severity": "重大", "title": "D-001 产品选型", "object_id": "DEC-001", "owner": "指导委员会"}
            ],
            "next_action": "推进 POC 证据汇总",
        },
    }


def make_suite(root: Path) -> Path:
    cost_dir = root / "05_成本与合同管理"
    cost_dir.mkdir(parents=True)
    cost = cost_dir / "成本与合同管理-20260918-V5.xlsx"
    book = Workbook()
    budget = book.active
    budget.title = "人工成本预算"
    budget["E3"] = 1382.27
    budget["A15"] = "合计"
    budget["D15"] = 4721
    budget["F15"] = 6525700
    detail = book.create_sheet("月度投入明细")
    detail["AP15"] = 4721
    detail["AQ15"] = 378
    detail["AT15"] = 378 / 4721
    book.save(cost)
    registry = root / "automation" / "config" / "current_workbooks.json"
    registry.parent.mkdir(parents=True)
    registry.write_text(json.dumps({
        "cost": cost.relative_to(root).as_posix(),
    }, ensure_ascii=False), encoding="utf-8")
    return cost


class ExtractCostOverviewTests(unittest.TestCase):
    def test_reads_budget_total_and_actual_completion(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            make_suite(root)
            overview = extract_cost_overview(root)
            self.assertEqual(overview["person_days"], 4721)
            self.assertEqual(overview["total_cost_wan"], "652.57")
            self.assertEqual(overview["average_day_rate"], "1,382.27")
            self.assertEqual(overview["actual_person_days"], 378)
            self.assertAlmostEqual(overview["completion_rate"], 378 / 4721, places=4)

    def test_returns_none_when_cost_workbook_missing(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            (root / "automation" / "config").mkdir(parents=True)
            (root / "automation" / "config" / "current_workbooks.json").write_text(
                json.dumps({}), encoding="utf-8")
            self.assertIsNone(extract_cost_overview(root))

    def test_returns_none_when_total_row_missing(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            cost = make_suite(root)
            book = Workbook()
            sheet = book.active
            sheet.title = "人工成本预算"
            sheet["A15"] = "别的"
            book.save(cost)
            self.assertIsNone(extract_cost_overview(root))


class PayloadCostTests(unittest.TestCase):
    def test_payload_includes_cost_overview(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            make_suite(root)
            overview = extract_cost_overview(root)
            payload = weekly_report_payload(make_report(), "合同系统", overview)
            self.assertEqual(payload["cost"]["person_days"], "4,721")
            self.assertEqual(payload["cost"]["total_cost_wan"], "652.57")
            self.assertEqual(payload["cost"]["actual_person_days"], "378")

    def test_payload_omits_cost_when_overview_none(self):
        payload = weekly_report_payload(make_report(), "合同系统", None)
        self.assertNotIn("cost", payload)


class DocxCostSectionTests(unittest.TestCase):
    def test_docx_contains_cost_section(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            make_suite(root)
            overview = extract_cost_overview(root)
            payload = weekly_report_payload(make_report(), "合同系统", overview)
            path = Path(folder) / "周报.docx"
            build_weekly_docx(payload, path)
            document = Document(path)
            headings = [p.text for p in document.paragraphs if p.style.name.startswith("Heading")]
            self.assertIn("五 成本投入概况", headings)
            table_text = "\n".join(
                cell.text for table in document.tables for row in table.rows for cell in row.cells
            )
            self.assertIn("4,721", table_text)
            self.assertIn("652.57", table_text)
            self.assertIn("1,382.27", table_text)
            self.assertIn("378", table_text)

    def test_docx_without_cost_keeps_original_five_sections(self):
        payload = weekly_report_payload(make_report(), "合同系统", None)
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "周报.docx"
            build_weekly_docx(payload, path)
            document = Document(path)
            headings = [p.text for p in document.paragraphs if p.style.name.startswith("Heading")]
            self.assertNotIn("五 成本投入概况", headings)
            self.assertIn("六 下周管理重点", headings)


if __name__ == "__main__":
    unittest.main()
