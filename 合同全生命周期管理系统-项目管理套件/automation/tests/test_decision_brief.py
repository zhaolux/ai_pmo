from __future__ import annotations

import json
import hashlib
import tempfile
import unittest
from datetime import date
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

from openpyxl import Workbook

from ai_pmo.decision_brief import (
    build_decision_brief, build_decision_payload, decision_brief_paths, extract_decision_facts,
)


MAPPING = [
    {"decision_id": "DEC-001", "package_id": "PKG-001", "budget_id": "BUD-002", "product": "CA及电子签章"},
    {"decision_id": "DEC-002", "package_id": "PKG-002", "budget_id": "BUD-003", "product": "电子文档"},
    {"decision_id": "DEC-003", "package_id": "PKG-003", "budget_id": "BUD-004", "product": "BI报表"},
]


def fixture_suite(root: Path, duplicate: bool = False, missing_package: bool = False) -> Path:
    communication = root / "08_沟通会议与报告"
    cost = root / "05_成本与合同管理"
    communication.mkdir(parents=True)
    cost.mkdir(parents=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "决策日志"
    ws.append(["决策日志"])
    ws.append([])
    ws.append([])
    ws.append(["决策ID", "决策主题", "类别", "提出日期", "决策截止日期", "背景/事实", "备选方案", "影响", "建议方案", "状态", "决策人", "决策日期", "决策结果", "证据/关联ID"])
    for i in range(1, 4):
        ws.append([f"DEC-{i:03d}", f"产品{i}POC选型", "技术/采购", None, date(2026, 9, 30), "待验证能力", "各候选厂商方案", "影响交付", "依据POC证据选型", "待决策", "项目治理委员会", None, None, "POC评分表"])
    if duplicate:
        ws.append(["DEC-001", "重复", None, None, date(2026, 9, 30)])
    wb.save(communication / "沟通会议与报告-20260918-V3.xlsx")
    wb = Workbook()
    ws = wb.active
    ws.title = "采购包计划"
    ws.append(["采购包计划"])
    ws.append([])
    ws.append([])
    ws.append(["采购包ID", "采购包名称", "采购类型", "范围与计价边界", "负责人", "估算完成日期", "计划签约日期", "计划金额(元)", "预算状态", "前置条件"])
    for i in range(1, 4):
        if missing_package and i == 2:
            continue
        ws.append([f"PKG-{i:03d}", f"产品{i}", "软件/服务", "许可及接口", "技术经理", date(2026, 9, 25), date(2026, 10, 15), None, "待估算", "POC后确认"])
    ws = wb.create_sheet("预算基线")
    ws.append(["预算基线"])
    ws.append([])
    ws.append([])
    ws.append(["预算ID", "成本类别", "成本项目", "一次性/经常性", "计量依据", "数量", "单价(元)", "初始预算(元)", "批准变更(元)", "当前预算(元)", "责任人", "估算状态", "估算依据/说明"])
    for i in range(2, 5):
        ws.append([f"BUD-{i:03d}", "软件成本", f"产品{i-1}", "一次性", "许可及接口", None, None, 0, 0, 0, "技术经理", "待估算", "等待POC和厂商报价"])
    wb.save(cost / "成本与合同管理-20260918-V5.xlsx")
    return root


class DecisionBriefTests(unittest.TestCase):
    def test_exact_joins_and_evidence(self):
        with tempfile.TemporaryDirectory() as temp:
            facts = extract_decision_facts(fixture_suite(Path(temp)), MAPPING, date(2026, 9, 19))
        self.assertEqual([(x["decision_id"], x["package_id"], x["budget_id"]) for x in facts], [
            ("DEC-001", "PKG-001", "BUD-002"),
            ("DEC-002", "PKG-002", "BUD-003"),
            ("DEC-003", "PKG-003", "BUD-004"),
        ])
        self.assertIsNone(facts[0]["planned_amount"])
        self.assertEqual(facts[0]["evidence"]["decision"]["row"], 5)

    def test_duplicate_or_missing_mapping_fails(self):
        with tempfile.TemporaryDirectory() as temp:
            suite = fixture_suite(Path(temp), duplicate=True)
            with self.assertRaisesRegex(ValueError, "重复决策ID: DEC-001"):
                extract_decision_facts(suite, MAPPING, date(2026, 9, 19))
        with tempfile.TemporaryDirectory() as temp:
            suite = fixture_suite(Path(temp), missing_package=True)
            with self.assertRaisesRegex(ValueError, "缺少采购包ID: PKG-002"):
                extract_decision_facts(suite, MAPPING, date(2026, 9, 19))

    def test_payload_marks_unverified_cost_and_scenarios(self):
        with tempfile.TemporaryDirectory() as temp:
            facts = extract_decision_facts(fixture_suite(Path(temp)), MAPPING, date(2026, 9, 19))
        payload = build_decision_payload(facts, None, date(2026, 9, 19))
        self.assertEqual(payload["approval_status"], "分析材料，非审批结论")
        self.assertEqual(payload["cards"][0]["cost_comparison"], "待估算")
        self.assertEqual(payload["cards"][0]["poc_score_status"], "未取得/未核验")
        self.assertEqual([x["title"] for x in payload["cards"][0]["scenarios"]], ["按现有计划选型", "补充POC后选型", "延期决策"])
        self.assertNotIn("推荐厂商", json.dumps(payload, ensure_ascii=False))

    def test_output_paths(self):
        json_path, docx_path = decision_brief_paths(Path("/tmp/suite"), date(2026, 9, 19))
        self.assertEqual(json_path.name, "POC选型决策简报-20260919-V1.json")
        self.assertEqual(docx_path.name, "POC选型决策简报-20260919-V1.docx")

    def test_builder_writes_three_cards_and_preserves_sources(self):
        from docx import Document
        with tempfile.TemporaryDirectory() as temp:
            suite = fixture_suite(Path(temp))
            sources = sorted([*suite.glob("08_沟通会议与报告/*.xlsx"), *suite.glob("05_成本与合同管理/*.xlsx")])
            before = {str(p): hashlib.sha256(p.read_bytes()).hexdigest() for p in sources}
            json_path, docx_path = build_decision_brief(suite, MAPPING, date(2026, 9, 19))
            payload = json.loads(json_path.read_text(encoding="utf-8"))
            text = "\n".join(p.text for p in Document(docx_path).paragraphs)
            self.assertEqual(len(payload["cards"]), 3)
            self.assertEqual(payload["approval_status"], "分析材料，非审批结论")
            for expected in ("DEC-001", "DEC-002", "DEC-003", "待估算", "未取得/未核验", "按现有计划选型", "补充POC后选型", "延期决策", "第5行"):
                self.assertIn(expected, text)
            self.assertEqual(before, {str(p): hashlib.sha256(p.read_bytes()).hexdigest() for p in sources})
            original_json = json_path.read_bytes()
            build_decision_brief(suite, MAPPING, date(2026, 9, 19))
            self.assertEqual(original_json, json_path.read_bytes())

    def test_missing_workbook_creates_no_output(self):
        with tempfile.TemporaryDirectory() as temp:
            suite = Path(temp)
            with self.assertRaises(FileNotFoundError):
                build_decision_brief(suite, MAPPING, date(2026, 9, 19))
            self.assertFalse((suite / "08_沟通会议与报告" / "决策简报").exists())

    def test_cli_dispatches_decision_brief(self):
        from ai_pmo.cli import execute
        paths = SimpleNamespace(suite=Path("/tmp/suite"), outputs=Path("/tmp/suite/outputs"))
        with patch("ai_pmo.cli.get_paths", return_value=paths), \
             patch("ai_pmo.cli.load_json", side_effect=lambda name: MAPPING if name == "poc_decision_mapping.json" else {"project_name": "合同系统"}), \
             patch("ai_pmo.cli.build_decision_brief", return_value=(Path("/tmp/a.json"), Path("/tmp/a.docx"))) as build, \
             patch("ai_pmo.cli._log"):
            self.assertEqual(execute("decision-brief", date(2026, 9, 19)), 0)
            build.assert_called_once()


if __name__ == "__main__":
    unittest.main()
