from __future__ import annotations

import json
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook

from scripts import sync_cost_to_ai_pmo as cost_sync


def make_suite(root: Path) -> tuple[Path, Path]:
    cost_dir = root / "05_成本与合同管理"
    target_dir = root / "11_AI_PMO中心"
    cost_dir.mkdir(parents=True)
    target_dir.mkdir(parents=True)
    cost = cost_dir / "成本与合同管理-20260920-V1.xlsx"
    book = Workbook()
    sheet = book.active
    sheet.title = "人工成本预算"
    sheet["A14"] = "合计"
    sheet["D14"] = 20
    sheet["F14"] = 24000
    budget = book.create_sheet("预算基线")
    budget["J5"] = 100
    budget["J6"] = 200
    budget["A20"] = "合计"
    budget["J20"] = 999  # '合计'行是求和结果，不应重复计入总预算
    book.save(cost)
    target = target_dir / "AI PMO中心-20260920-V1.xlsx"
    target.write_bytes(b"human-edited-workbook")
    registry = root / "automation" / "config" / "current_workbooks.json"
    registry.parent.mkdir(parents=True)
    registry.write_text(json.dumps({
        "cost": cost.relative_to(root).as_posix(),
        "ai_pmo": target.relative_to(root).as_posix(),
    }, ensure_ascii=False), encoding="utf-8")
    return cost, target


class CostSyncSafetyTests(unittest.TestCase):
    def test_preview_shows_cost_values_and_does_not_write(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            source, target = make_suite(root)

            preview = cost_sync.preview_sync(root)

            self.assertEqual(preview["source"], source)
            self.assertEqual(preview["target"], target)
            self.assertEqual(preview["metrics"].average_day_rate, 1200)
            self.assertEqual(preview["metrics"].total_budget, 300)
            self.assertEqual(target.read_bytes(), b"human-edited-workbook")
            self.assertFalse((target.parent / "_archive").exists())

    def test_preview_ignores_newer_unregistered_cost_workbook(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            source, _ = make_suite(root)
            newer = source.parent / "成本与合同管理-20260921-V1.xlsx"
            newer.write_bytes(b"unapproved")

            self.assertEqual(cost_sync.preview_sync(root)["source"], source)

    def test_default_main_does_not_apply(self):
        preview = {
            "source": Path("/tmp/cost.xlsx"), "target": Path("/tmp/pmo.xlsx"),
            "metrics": cost_sync.CostMetrics(20, 24000, 1200, 300),
        }
        with patch("sys.argv", ["sync_cost_to_ai_pmo.py", "--suite", "/tmp/suite"]), patch.object(
            cost_sync, "preview_sync", return_value=preview
        ), patch.object(cost_sync, "sync") as apply:
            cost_sync.main()
            apply.assert_not_called()

    def test_total_row_located_by_label_not_fixed_row(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            source, _ = make_suite(root)
            book = __import__("openpyxl").load_workbook(source)
            sheet = book["人工成本预算"]
            sheet.insert_rows(14, 1)
            sheet["A14"] = "中级测试人员"
            sheet["A15"] = "合计"
            sheet["D15"] = 21
            sheet["F15"] = 25200
            book.save(source)

            metrics = cost_sync.extract_cost_metrics(source)

            self.assertEqual(metrics.person_days, 21)
            self.assertEqual(metrics.total_cost, 25200)
            self.assertEqual(metrics.average_day_rate, 1200)

    def test_apply_rejects_target_changed_after_preview(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            _, old = make_suite(root)
            newer = old.parent / "AI PMO中心-20260920-V2.xlsx"
            newer.write_bytes(b"newer")
            registry = root / "automation" / "config" / "current_workbooks.json"
            values = json.loads(registry.read_text(encoding="utf-8"))
            values["ai_pmo"] = newer.relative_to(root).as_posix()
            registry.write_text(json.dumps(values, ensure_ascii=False), encoding="utf-8")

            with self.assertRaisesRegex(RuntimeError, "预览后已变化"):
                cost_sync.sync(root, expected_target=old)

            self.assertEqual(old.read_bytes(), b"human-edited-workbook")
            self.assertEqual(newer.read_bytes(), b"newer")

    def test_apply_backs_up_before_only_cost_cells_are_sent(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            source, target = make_suite(root)
            captured = {}

            def fake_run(command, *, check):
                captured["commands"] = json.loads(Path(command[4]).read_text(encoding="utf-8"))
                self.assertTrue(any((target.parent / "_archive" / "自动刷新前备份").glob("*.xlsx")))

            with patch.object(cost_sync.shutil, "which", return_value="officecli"), patch.object(
                cost_sync.subprocess, "run", side_effect=fake_run
            ):
                result_source, result_target, metrics, backup = cost_sync.sync(root, expected_target=target)

            self.assertEqual((result_source, result_target), (source, target))
            self.assertEqual(metrics.average_day_rate, 1200)
            self.assertEqual(backup.read_bytes(), b"human-edited-workbook")
            self.assertEqual(target.read_bytes(), b"human-edited-workbook")
            paths = {command["path"] for command in captured["commands"]}
            self.assertEqual(paths, {
                "/项目驾驶舱/B9", "/项目驾驶舱/C9", "/项目驾驶舱/D9", "/项目驾驶舱/I9",
                "/_数据接口/B16", "/_数据接口/B17", "/_数据接口/B18",
                "/_数据接口/B19", "/_数据接口/B20", "/_数据接口/B21",
                "/_数据接口/A43", "/_数据接口/B43",
            })
            budget_commands = {
                command["path"]: command["props"]
                for command in captured["commands"]
                if command["path"] in {"/项目驾驶舱/I9", "/_数据接口/A43", "/_数据接口/B43"}
            }
            self.assertEqual(budget_commands["/项目驾驶舱/I9"]["formula"], "'_数据接口'!B43")
            self.assertEqual(budget_commands["/_数据接口/A43"]["value"], "已量化总预算")
            self.assertEqual(budget_commands["/_数据接口/B43"], {"value": 300, "type": "number"})


if __name__ == "__main__":
    unittest.main()
