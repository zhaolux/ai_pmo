from __future__ import annotations

import json
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook

from ai_pmo.source_ownership import audit_source_ownership


def save_book(path: Path, sheet_name: str, header_row: int, header: str) -> None:
    book = Workbook()
    sheet = book.active
    sheet.title = sheet_name
    for _ in range(header_row - 1):
        sheet.append([])
    sheet.append([header, "名称"])
    sheet.append(["ID-001", "记录"])
    path.parent.mkdir(parents=True, exist_ok=True)
    book.save(path)


class SourceOwnershipTests(unittest.TestCase):
    def make_suite(self, root: Path) -> Path:
        suite = root / "suite"
        plan = suite / "02_计划与进度管理" / "计划与进度管理-20260923-V1.xlsx"
        risk = suite / "06_风险问题与变更" / "风险管理-20260923-V1.xlsx"
        save_book(plan, "项目总控台账", 3, "任务ID")
        save_book(risk, "风险登记册", 8, "风险ID")
        config = suite / "automation" / "config"
        config.mkdir(parents=True)
        (config / "current_workbooks.json").write_text(json.dumps({
            "plan": plan.relative_to(suite).as_posix(),
            "risk": risk.relative_to(suite).as_posix(),
        }, ensure_ascii=False), encoding="utf-8")
        (config / "data_ownership.json").write_text(json.dumps({"objects": [
            {"object": "task", "workbook": "plan", "sheet": "项目总控台账", "header_row": 3, "key_header": "任务ID"},
            {"object": "risk", "workbook": "risk", "sheet": "风险登记册", "header_row": 8, "key_header": "风险ID"},
        ]}, ensure_ascii=False), encoding="utf-8")
        return suite

    def test_valid_registry_has_one_owner_per_object(self):
        with tempfile.TemporaryDirectory() as directory:
            result = audit_source_ownership(self.make_suite(Path(directory)))
        self.assertEqual(result["counts"], {"objects": 2, "records": 2})
        self.assertEqual(result["errors"], [])

    def test_duplicate_object_owner_is_rejected(self):
        with tempfile.TemporaryDirectory() as directory:
            suite = self.make_suite(Path(directory))
            path = suite / "automation" / "config" / "data_ownership.json"
            payload = json.loads(path.read_text(encoding="utf-8"))
            payload["objects"].append(dict(payload["objects"][0]))
            path.write_text(json.dumps(payload, ensure_ascii=False), encoding="utf-8")
            result = audit_source_ownership(suite)
        self.assertTrue(any("重复主源" in item for item in result["errors"]))

    def test_missing_sheet_and_key_header_are_reported(self):
        with tempfile.TemporaryDirectory() as directory:
            suite = self.make_suite(Path(directory))
            path = suite / "automation" / "config" / "data_ownership.json"
            payload = json.loads(path.read_text(encoding="utf-8"))
            payload["objects"][0]["sheet"] = "不存在"
            payload["objects"][1]["key_header"] = "错误字段"
            path.write_text(json.dumps(payload, ensure_ascii=False), encoding="utf-8")
            result = audit_source_ownership(suite)
        self.assertTrue(any("Sheet不存在" in item for item in result["errors"]))
        self.assertTrue(any("主键表头" in item for item in result["errors"]))

    def test_optional_data_range_excludes_second_table_in_same_sheet(self):
        with tempfile.TemporaryDirectory() as directory:
            suite = self.make_suite(Path(directory))
            plan = suite / "02_计划与进度管理" / "计划与进度管理-20260923-V1.xlsx"
            book = __import__("openpyxl").load_workbook(plan)
            sheet = book["项目总控台账"]
            sheet.cell(5, 1, "任务ID")
            sheet.cell(6, 1, "ID-001")
            book.save(plan)
            path = suite / "automation" / "config" / "data_ownership.json"
            payload = json.loads(path.read_text(encoding="utf-8"))
            payload["objects"][0].update({"data_start_row": 4, "data_end_row": 4})
            path.write_text(json.dumps(payload, ensure_ascii=False), encoding="utf-8")
            result = audit_source_ownership(suite)
        task = next(item for item in result["objects"] if item["object"] == "task")
        self.assertEqual(task["records"], 1)
        self.assertFalse(any("task 主键重复" in item for item in result["errors"]))


if __name__ == "__main__":
    unittest.main()
