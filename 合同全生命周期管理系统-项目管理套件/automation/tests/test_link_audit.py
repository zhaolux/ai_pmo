from __future__ import annotations

import unittest
import tempfile
import json
from pathlib import Path

from ai_pmo.link_audit import audit_interface_relations, audit_suite_links, compare_snapshots, resource_task_ids
from ai_pmo.current_workbooks import resolve_current


class LinkAuditTests(unittest.TestCase):
    def test_snapshot_detects_changed_field_and_duplicate_id(self):
        source = [(5, ("IF-001", "ERP", "P0")), (6, ("IF-001", "ERP", "P0"))]
        copy = [(5, ("IF-001", "ERP", "P1"))]
        findings = compare_snapshots(source, copy, "接口清单", ("清单ID", "系统", "优先级"))
        self.assertIn("接口清单:源表重复ID IF-001 行5/6", findings)
        self.assertIn("接口清单:IF-001 优先级不一致 源行5/副本行5", findings)

    def test_relations_find_unknown_parent_ids(self):
        specs = [(5, ("SPEC-001", "IF-404"))]
        implementations = [(5, ("IMP-001", "SPEC-404", "IF-001", "T404"))]
        findings = audit_interface_relations({"IF-001"}, specs, implementations, {"T001"})
        self.assertIn("P0规格:SPEC-001 引用不存在的清单ID IF-404 行5", findings)
        self.assertIn("实施任务:IMP-001 引用不存在的规格ID SPEC-404 行5", findings)
        self.assertIn("实施任务:IMP-001 引用不存在的总控任务 T404 行5", findings)

    def test_relations_reject_spec_interface_mismatch(self):
        specs = [(5, ("SPEC-001", "IF-001"))]
        implementations = [(5, ("IMP-001", "SPEC-001", "IF-002", "T001"))]
        findings = audit_interface_relations({"IF-001", "IF-002"}, specs, implementations, {"T001"})
        self.assertIn("实施任务:IMP-001 规格SPEC-001所属接口IF-001与任务接口IF-002不一致 行5", findings)

    def test_resource_task_ids_extracts_only_task_references(self):
        self.assertEqual(resource_task_ids(["T001 方案；T002 评审", "无", "T001 方案"]), {"T001", "T002"})

    def test_scope_and_integration_have_explicit_current_versions(self):
        with tempfile.TemporaryDirectory() as folder:
            suite = Path(folder)
            paths = {
                "scope": "03_需求与范围管理/需求与范围管理-20260918-V7.xlsx",
                "integration": "04_系统集成管理/系统集成管理-20260918-V1.xlsx",
            }
            config = suite / "automation/config"
            config.mkdir(parents=True)
            (config / "current_workbooks.json").write_text(json.dumps(paths), encoding="utf-8")
            for relative in paths.values():
                target = suite / relative
                target.parent.mkdir(parents=True)
                target.write_bytes(b"fixture")
            for key, relative in paths.items():
                self.assertEqual(resolve_current(suite, key), suite / relative)

    def test_audit_suite_reports_stale_snapshot_and_unplanned_task(self):
        from openpyxl import Workbook

        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            scope = Workbook()
            s = scope.active
            s.title = "接口清单"
            s.append([]); s.append([]); s.append([])
            s.append(["清单ID", "优先级"])
            s.append(["IF-001", "P0"])
            spec = scope.create_sheet("P0接口规格台账")
            for _ in range(3): spec.append([])
            spec.append(["规格ID", "清单ID"])
            spec.append(["SPEC-001", "IF-001"])
            scope_path = root / "scope.xlsx"
            scope.save(scope_path)

            integration = Workbook()
            i = integration.active
            i.title = "接口清单"
            for _ in range(3): i.append([])
            i.append(["清单ID", "优先级"])
            i.append(["IF-001", "P1"])
            spec_copy = integration.create_sheet("P0接口规格")
            for _ in range(3): spec_copy.append([])
            spec_copy.append(["规格ID", "清单ID"])
            spec_copy.append(["SPEC-001", "IF-001"])
            impl = integration.create_sheet("接口实施计划")
            for _ in range(3): impl.append([])
            impl.append(["实施任务ID", "规格ID", "清单ID", "关联总控任务"])
            impl.append(["IMP-001", "SPEC-001", "IF-001", "T001"])
            integration_path = root / "integration.xlsx"
            integration.save(integration_path)

            plan = Workbook()
            t = plan.active
            t.title = "项目总控台账"
            t.append([]); t.append([]); t.append(["任务ID"]); t.append(["T001"])
            resource = plan.create_sheet("资源计划")
            for _ in range(37): resource.append([])
            resource.append(["人员", "周开始", "周结束", "是否有任务", "任务数", "周计划工时", "负荷率", "负荷判断", "任务清单"])
            resource.append(["开发-01", None, None, "无", 0, 0, None, None, "无"])
            plan_path = root / "plan.xlsx"
            plan.save(plan_path)

            result = audit_suite_links(scope_path, integration_path, plan_path)
            self.assertTrue(any("优先级不一致" in issue for issue in result["errors"]))
            self.assertTrue(any("T001" in issue for issue in result["warnings"]))
            self.assertEqual(result["counts"]["实施任务"], 1)


if __name__ == "__main__":
    unittest.main()
