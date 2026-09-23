"""Read-only cross-check of contract migration planning gates.

The four workbook paths are explicit inputs. No workbook is written or approved here.
"""

from __future__ import annotations

import argparse
import json
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook


def _record(path: Path, sheet: str, record_id: str, errors: list[str]):
    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
        worksheet = workbook[sheet]
    except (OSError, KeyError) as exc:
        errors.append(f'{path.name} / {sheet}: {exc}')
        return None
    matches = []
    for row_number, cells in enumerate(worksheet.iter_rows(values_only=True), 1):
        if cells and cells[0] == record_id:
            matches.append((row_number, cells))
    workbook.close()
    if len(matches) != 1:
        errors.append(f'{path.name} / {sheet}: {record_id} 应恰有一条，实际 {len(matches)} 条')
        return None
    row_number, cells = matches[0]
    return {'file': str(path), 'sheet': sheet, 'row': row_number, 'values': cells}


def _iso(value):
    return value.date().isoformat() if isinstance(value, datetime) else value.isoformat() if isinstance(value, date) else None


SOURCE_REQUIRED_FIELDS = {
    3: '实际系统/库名称',
    4: '是否独立系统',
    5: '数据责任人',
    6: '合同/记录量',
    7: '附件量',
    8: '抽取方式',
    9: '历史保留边界',
    10: '证据/确认人',
}
SOURCE_PENDING_VALUES = {'', '待盘点', '待确认', '待指定', '待核实', '待补充'}


def _source_inventory(path: Path, errors: list[str]) -> dict:
    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
        worksheet = workbook['来源系统盘点']
        clues = [(number, cells) for number, cells in enumerate(worksheet.iter_rows(values_only=True), 1)
                 if cells and isinstance(cells[0], str) and cells[0].startswith('SRC-')]
        workbook.close()
    except (OSError, KeyError) as exc:
        errors.append(f'{path.name} / 来源系统盘点: {exc}')
        clues = []
    records = []
    for number, cells in clues:
        missing = [name for index, name in SOURCE_REQUIRED_FIELDS.items()
                   if index >= len(cells) or str(cells[index] if cells[index] is not None else '').strip() in SOURCE_PENDING_VALUES]
        if len(cells) <= 11 or cells[11] != '已确认':
            missing.append('盘点状态')
        records.append({
            'source': {'file': str(path), 'sheet': '来源系统盘点', 'row': number, 'id': cells[0]},
            'missing_fields': missing,
        })
    return {'clue_count': len(clues), 'records': records,
            'fields_complete': bool(clues) and all(not item['missing_fields'] for item in records)}


def planned_sequence_issue(rehearsal_end: str | None, migration_dates: list[str | None], cut_start: str | None) -> str | None:
    if not rehearsal_end or not migration_dates or not all(migration_dates) or not cut_start:
        return '迁移阶段计划日期不完整，需核实'
    if rehearsal_end > min(migration_dates):
        return 'T087 演练完成日晚于 MIG 最早计划执行日，需核实阶段定义和依赖'
    if max(migration_dates) > cut_start:
        return 'MIG 计划执行日晚于 CUT-005 生产切换开始日，需核实阶段定义和依赖'
    return None


def _quality_coverage(path: Path, migrations: dict, errors: list[str]) -> dict:
    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
        worksheet = workbook['数据质量与迁移']
        quality_rows = [(number, cells) for number, cells in enumerate(worksheet.iter_rows(values_only=True), 1)
                        if len(cells) > 1 and cells[0] not in (None, '数据对象') and cells[1] not in (None, '来源系统')]
        workbook.close()
    except (OSError, KeyError) as exc:
        errors.append(f'{path.name} / 数据质量与迁移: {exc}')
        return {'planned_rule_coverage': None, 'verified_count': 0, 'records': []}

    by_source = {}
    for number, cells in quality_rows:
        by_source.setdefault(cells[1], []).append((number, cells))
    records = []
    covered = verified = 0
    for migration_id, record in migrations.items():
        if record is None:
            continue
        cells = record['values']
        source = cells[2] if len(cells) > 2 else None
        matches = by_source.get(source, []) if source else []
        if len(matches) != 1:
            errors.append(f'{migration_id}: 质量台账来源“{source}”匹配 {len(matches)} 条，无法唯一关联')
            continue
        number, quality = matches[0]
        rule_covered = all(len(quality) > column and quality[column] not in (None, '') for column in (4, 5, 6))
        completed = (len(quality) > 11 and quality[8] not in (None, '')
                     and quality[10] in ('已完成', '已验证', '已关闭') and quality[11] not in (None, ''))
        covered += int(rule_covered)
        verified += int(completed)
        records.append({
            'migration_id': migration_id,
            'migration_source': {'file': record['file'], 'sheet': record['sheet'], 'row': record['row']},
            'quality_source': {'file': str(path), 'sheet': '数据质量与迁移', 'row': number},
            'data_object': quality[0],
            'rule_covered': rule_covered,
            'verification_status': quality[10] if len(quality) > 10 else None,
            'verified_with_evidence': completed,
        })
    return {'planned_rule_coverage': f'{covered}/{len(migrations)}', 'verified_count': verified, 'records': records}


def audit_migration_gates(gate: Path, plan: Path, launch: Path, migration: Path,
                          quality: Path | None = None, risk: Path | None = None) -> dict:
    """Return evidence-linked facts, missing inputs and structural errors; never infer approval."""
    files = tuple(Path(path) for path in (gate, plan, launch, migration))
    errors: list[str] = []
    found = {
        'G15': _record(files[0], '启动门禁清单', 'G15', errors),
        'T087': _record(files[1], '项目总控台账', 'T087', errors),
        'T092': _record(files[1], '项目总控台账', 'T092', errors),
        'RDY-005': _record(files[2], '上线准备清单', 'RDY-005', errors),
        'CUT-005': _record(files[2], '切换执行计划', 'CUT-005', errors),
    }
    for number in range(1, 7):
        key = f'MIG-{number:03d}'
        found[key] = _record(files[2], '数据迁移与核验', key, errors)

    source_inventory = _source_inventory(files[3], errors)

    actions = []
    gate_record = found['G15']
    if gate_record and (len(gate_record['values']) < 8 or gate_record['values'][7] != '已通过'):
        actions.append('G15 待确认')
    if not source_inventory['fields_complete']:
        actions.append('来源系统待盘点')
    ready = found['RDY-005']
    if ready and (len(ready['values']) < 11 or ready['values'][10] != '已就绪'):
        actions.append('RDY-005 未就绪')

    def item(key, status_col, start_col, end_col=None):
        record = found[key]
        if record is None:
            return None
        cells = record['values']
        return {
            'source': {'file': record['file'], 'sheet': record['sheet'], 'row': record['row'], 'id': key},
            'status': cells[status_col] if len(cells) > status_col else None,
            'planned_start': _iso(cells[start_col]) if len(cells) > start_col else None,
            'planned_end': _iso(cells[end_col]) if end_col is not None and len(cells) > end_col else None,
        }

    records = {
        'G15': item('G15', 7, 6),
        'T087': item('T087', 6, 8, 9),
        'T092': item('T092', 6, 8, 9),
        'RDY-005': item('RDY-005', 10, 5),
        'CUT-005': item('CUT-005', 12, 6, 7),
    }
    records.update({f'MIG-{n:03d}': item(f'MIG-{n:03d}', 12, 6) for n in range(1, 7)})
    if not errors:
        sequence_issue = planned_sequence_issue(
            records['T087']['planned_end'],
            [records[f'MIG-{n:03d}']['planned_start'] for n in range(1, 7)],
            records['CUT-005']['planned_start'],
        )
        if sequence_issue:
            actions.append(sequence_issue)
    quality_result = None
    if quality is not None:
        quality_result = _quality_coverage(Path(quality), {f'MIG-{n:03d}': found[f'MIG-{n:03d}'] for n in range(1, 7)}, errors)
        if quality_result['planned_rule_coverage'] != '6/6':
            actions.append('迁移质量规则覆盖待核对')
        if quality_result['verified_count'] < 6:
            actions.append(f"迁移质量验证未完成 {6 - quality_result['verified_count']} 项")
    risk_result = None
    if risk is not None:
        risk_record = _record(Path(risk), '风险登记册', 'RSK-007', errors)
        if risk_record is not None:
            risk_cells = risk_record['values']
            value = lambda index: risk_cells[index] if len(risk_cells) > index else None
            risk_result = {
                'source': {'file': risk_record['file'], 'sheet': risk_record['sheet'], 'row': risk_record['row'], 'id': 'RSK-007'},
                'name': value(1),
                'owner': value(22),
                'response_due': _iso(value(27)),
                'status': value(28),
                'residual_level': value(33),
                'closure_evidence_text': value(40),
            }
            if risk_result['status'] != '已关闭':
                actions.append(f"RSK-007 {risk_result['status'] or '状态待确认'}")
            else:
                actions.append('RSK-007 关闭证据需人工核实')
    return {
        'planned_sequence': '演练 → 计划执行 → 生产切换',
        'records': records,
        'source_clue_count': source_inventory['clue_count'],
        'source_inventory': source_inventory,
        'confirmed_system_count': None,
        'quality': quality_result,
        'risk': risk_result,
        'actions': actions,
        'errors': errors,
        'note': 'MIG 的日期列仅为计划执行日；线索字段齐全也不能证明已穷尽独立系统，线索行数不是独立系统数量。状态是源台账快照，不代表审批。',
    }


def main() -> int:
    parser = argparse.ArgumentParser(description='只读核对合同数据迁移门禁，不修改工作簿')
    for name in ('gate', 'plan', 'launch', 'migration'):
        parser.add_argument(f'--{name}', type=Path, required=True)
    parser.add_argument('--quality', type=Path, help='可选：质量测试与验收工作簿，核对迁移规则覆盖和验证证据')
    parser.add_argument('--risk', type=Path, help='可选：风险管理工作簿，核对 RSK-007 责任、状态和应对证据')
    args = parser.parse_args()
    result = audit_migration_gates(args.gate, args.plan, args.launch, args.migration, args.quality, args.risk)
    print(json.dumps(result, ensure_ascii=False, indent=2))
    return 1 if result['errors'] else 0


if __name__ == '__main__':
    raise SystemExit(main())
