from __future__ import annotations

import json
import os
import tempfile
from datetime import date, datetime
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from openpyxl import load_workbook


def _latest(directory: Path, prefix: str) -> Path:
    matches = [p for p in directory.glob(f"{prefix}*.xlsx") if not p.name.startswith("~$")]
    if not matches:
        raise FileNotFoundError(f"缺少必需工作簿: {directory}/{prefix}*.xlsx")
    return max(matches, key=lambda p: (p.stat().st_mtime, p.name))


def _records(path: Path, sheet: str, id_header: str, required: tuple[str, ...]) -> dict[str, tuple[dict, int]]:
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        if sheet not in wb.sheetnames:
            raise ValueError(f"缺少工作表: {path.name}/{sheet}")
        ws = wb[sheet]
        headings = [cell.value for cell in ws[4]]
        absent = [name for name in required if name not in headings]
        if absent:
            raise ValueError(f"缺少字段: {path.name}/{sheet}: {', '.join(absent)}")
        index = {name: headings.index(name) for name in required}
        result = {}
        for row_num, values in enumerate(ws.iter_rows(min_row=5, values_only=True), 5):
            key = values[index[id_header]] if index[id_header] < len(values) else None
            if key is None or str(key).strip() == "":
                continue
            key = str(key).strip()
            if key in result:
                raise ValueError(f"重复{id_header}: {key}")
            result[key] = ({name: values[col] if col < len(values) else None for name, col in index.items()}, row_num)
        return result
    finally:
        wb.close()


def _date(value: object) -> str | None:
    if isinstance(value, (date, datetime)):
        return value.date().isoformat() if isinstance(value, datetime) else value.isoformat()
    return str(value)[:10] if value else None


def _evidence(path: Path, sheet: str, record_id: str, row: int) -> dict:
    return {"file": path.name, "sheet": sheet, "record_id": record_id, "row": row}


def extract_decision_facts(suite: Path, mapping: list[dict], as_of: date) -> list[dict]:
    communication = _latest(suite / "08_沟通会议与报告", "沟通会议与报告")
    cost = _latest(suite / "05_成本与合同管理", "成本与合同管理")
    decisions = _records(communication, "决策日志", "决策ID", (
        "决策ID", "决策主题", "决策截止日期", "背景/事实", "备选方案", "影响",
        "建议方案", "状态", "决策人", "证据/关联ID",
    ))
    packages = _records(cost, "采购包计划", "采购包ID", (
        "采购包ID", "采购包名称", "范围与计价边界", "负责人", "估算完成日期",
        "计划签约日期", "计划金额(元)", "预算状态", "前置条件",
    ))
    budgets = _records(cost, "预算基线", "预算ID", (
        "预算ID", "成本项目", "当前预算(元)", "估算状态", "估算依据/说明",
    ))
    if len({m["decision_id"] for m in mapping}) != len(mapping):
        raise ValueError("决策映射含重复决策ID")
    facts = []
    for item in mapping:
        decision_id, package_id, budget_id = (item[name] for name in ("decision_id", "package_id", "budget_id"))
        for identifier, source, label in ((decision_id, decisions, "决策ID"), (package_id, packages, "采购包ID"), (budget_id, budgets, "预算ID")):
            if identifier not in source:
                raise ValueError(f"缺少{label}: {identifier}")
        decision, decision_row = decisions[decision_id]
        package, package_row = packages[package_id]
        budget, budget_row = budgets[budget_id]
        facts.append({
            "product": item["product"], "decision_id": decision_id,
            "decision_title": decision["决策主题"], "decision_due": _date(decision["决策截止日期"]),
            "decision_status": decision["状态"], "decision_owner": decision["决策人"],
            "background": decision["背景/事实"], "options": decision["备选方案"],
            "impact": decision["影响"], "source_recommendation": decision["建议方案"],
            "poc_reference": decision["证据/关联ID"],
            "package_id": package_id, "package_name": package["采购包名称"],
            "pricing_scope": package["范围与计价边界"], "package_owner": package["负责人"],
            "estimate_due": _date(package["估算完成日期"]),
            "contract_due": _date(package["计划签约日期"]),
            "planned_amount": package["计划金额(元)"], "package_budget_status": package["预算状态"],
            "package_prerequisite": package["前置条件"],
            "budget_id": budget_id, "budget_item": budget["成本项目"],
            "budget_amount": budget["当前预算(元)"], "budget_status": budget["估算状态"],
            "budget_basis": budget["估算依据/说明"],
            "evidence": {
                "decision": _evidence(communication, "决策日志", decision_id, decision_row),
                "package": _evidence(cost, "采购包计划", package_id, package_row),
                "budget": _evidence(cost, "预算基线", budget_id, budget_row),
            },
        })
    return facts


def build_decision_payload(facts: list[dict], agent_report: dict | None, as_of: date) -> dict:
    cards = []
    for fact in facts:
        alerts = []
        if agent_report is not None and agent_report.get("data_date") == as_of.isoformat():
            alerts = [
                {"title": f["title"], "detail": f["detail"], "evidence": f.get("evidence", [])}
                for f in agent_report.get("findings", []) if f.get("object_id") == fact["decision_id"]
            ]
        missing = ["POC评分表实际文件与评分结果未取得/未核验", "候选厂商报价未取得/未核验"]
        if fact["planned_amount"] is None or fact["budget_status"] == "待估算":
            missing.append(f"{fact['package_id']}采购估算及{fact['budget_id']}预算基线待量化，最迟{fact['estimate_due']}")
        cards.append({
            "decision_id": fact["decision_id"], "product": fact["product"],
            "title": fact["decision_title"], "due_date": fact["decision_due"],
            "status": fact["decision_status"], "decision_owner": fact["decision_owner"],
            "facts": fact, "poc_score_status": "未取得/未核验",
            "cost_comparison": "待估算" if fact["planned_amount"] is None or fact["budget_status"] == "待估算" else "可比较，须人工确认口径",
            "missing_evidence": missing, "agent_alerts": alerts,
            "scenarios": [
                {"title": "按现有计划选型", "type": "分析假设", "impact": "如POC评分、报价与合规证据及时齐备，可维持选型节点；证据不足时不得据此批准采购。"},
                {"title": "补充POC后选型", "type": "分析假设", "impact": "可提高能力和成本判断的可靠性，但需重新确认选型与采购时间。"},
                {"title": "延期决策", "type": "分析假设", "impact": "可能影响产品集成、签约和项目启动准备；具体进度影响须依据依赖计划测算。"},
            ],
            "manual_decision": "",
        })
    return {"data_date": as_of.isoformat(), "approval_status": "分析材料，非审批结论", "cards": cards}


def decision_brief_paths(suite: Path, as_of: date) -> tuple[Path, Path]:
    directory = suite / "08_沟通会议与报告" / "决策简报"
    stem = f"POC选型决策简报-{as_of:%Y%m%d}-V1"
    return directory / f"{stem}.json", directory / f"{stem}.docx"


def _render_word(payload: dict, path: Path) -> None:
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Cm(1.8)
    sec.bottom_margin = Cm(1.8)
    sec.left_margin = Cm(2.2)
    sec.right_margin = Cm(2.2)
    for name, size in (("Normal", 10), ("Title", 18), ("Heading 1", 13), ("Heading 2", 11)):
        style = doc.styles[name]
        style.font.name = "SimSong"
        style._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "SimSong")
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
    title_properties = doc.styles["Title"]._element.get_or_add_pPr()
    title_border = title_properties.find(qn("w:pBdr"))
    if title_border is not None:
        title_properties.remove(title_border)
    doc.add_paragraph("合同系统外部产品 POC 选型决策简报", style="Title")
    doc.add_paragraph(f"数据截止日 {payload['data_date']}。本材料供项目治理委员会核验证据与讨论选型，不构成审批结论。")
    doc.add_paragraph("当前三项产品均待决策。POC评分结果和候选厂商报价尚未核验，采购金额待估算；请在决定选型前补齐这些证据。")
    for card_index, card in enumerate(payload["cards"]):
        if card_index:
            doc.add_page_break()
        fact = card["facts"]
        doc.add_heading(f"{card['decision_id']} {card['product']}", level=1)
        doc.add_paragraph(f"{card['title']}；截止日 {card['due_date']}；状态 {card['status']}；决策人 {card['decision_owner']}。")
        doc.add_heading("已记录的事实与来源", level=2)
        for label, value, source in (
            ("待验证能力", fact["background"], "decision"),
            ("备选路径", fact["options"], "decision"),
            ("影响", fact["impact"], "decision"),
            ("采购范围", fact["pricing_scope"], "package"),
            ("估算状态", f"{fact['package_budget_status']}；估算计划日 {fact['estimate_due']}", "package"),
            ("预算状态", f"{fact['budget_status']}；成本比较{card['cost_comparison']}", "budget"),
        ):
            e = fact["evidence"][source]
            doc.add_paragraph(f"{label}：{value}。来源：{e['file']}，工作表{e['sheet']}，记录{e['record_id']}，第{e['row']}行。")
        doc.add_heading("待补证据", level=2)
        for item in card["missing_evidence"]:
            doc.add_paragraph(item, style="List Bullet")
        doc.add_paragraph(f"POC评分状态：{card['poc_score_status']}；台账引用“{fact['poc_reference']}”不代表文件已取得。")
        doc.add_heading("供讨论的情景假设", level=2)
        for item in card["scenarios"]:
            doc.add_paragraph(f"{item['title']}（{item['type']}）：{item['impact']}")
        doc.add_paragraph("人工决策结论：____________________；决策日期：____________")
    doc.add_paragraph("分析材料，非审批结论。审批结果请由授权人录入源决策日志。")
    doc.save(path)


def build_decision_brief(suite: Path, mapping: list[dict], as_of: date, agent_report: dict | None = None) -> tuple[Path, Path]:
    facts = extract_decision_facts(suite, mapping, as_of)
    payload = build_decision_payload(facts, agent_report, as_of)
    json_path, docx_path = decision_brief_paths(suite, as_of)
    json_path.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(dir=json_path.parent, prefix=".decision-brief-") as temp:
        draft_json = Path(temp) / json_path.name
        draft_docx = Path(temp) / docx_path.name
        draft_json.write_text(json.dumps(payload, ensure_ascii=False, indent=2, sort_keys=True), encoding="utf-8")
        _render_word(payload, draft_docx)
        os.replace(draft_json, json_path)
        os.replace(draft_docx, docx_path)
    return json_path, docx_path
