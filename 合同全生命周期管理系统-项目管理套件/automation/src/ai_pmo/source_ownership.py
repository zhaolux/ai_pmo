from __future__ import annotations

import json
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook

from .current_workbooks import resolve_current


def _load_contract(suite: Path) -> list[dict]:
    path = suite / "automation" / "config" / "data_ownership.json"
    if not path.is_file():
        raise FileNotFoundError(f"缺少数据主源注册表：{path}")
    payload = json.loads(path.read_text(encoding="utf-8"))
    objects = payload.get("objects") if isinstance(payload, dict) else None
    if not isinstance(objects, list) or not objects:
        raise ValueError("数据主源注册表没有objects定义")
    return objects


def audit_source_ownership(suite: Path) -> dict:
    contracts = _load_contract(suite)
    errors: list[str] = []
    details: list[dict] = []
    names = [item.get("object") for item in contracts if isinstance(item, dict)]
    for name, count in Counter(names).items():
        if not name:
            errors.append("数据主源存在空对象名称")
        elif count > 1:
            errors.append(f"对象存在重复主源：{name}（{count}处）")

    records = 0
    for contract in contracts:
        if not isinstance(contract, dict):
            errors.append("数据主源条目不是对象")
            continue
        name = str(contract.get("object") or "")
        workbook_key = str(contract.get("workbook") or "")
        sheet_name = str(contract.get("sheet") or "")
        key_header = str(contract.get("key_header") or "")
        header_row = contract.get("header_row")
        data_start_row = contract.get("data_start_row", header_row + 1 if isinstance(header_row, int) else None)
        data_end_row = contract.get("data_end_row")
        if not all((name, workbook_key, sheet_name, key_header)) or not isinstance(header_row, int) or header_row < 1:
            errors.append(f"数据主源定义不完整：{name or '未命名对象'}")
            continue
        if not isinstance(data_start_row, int) or data_start_row <= header_row:
            errors.append(f"{name} 数据起始行无效")
            continue
        if data_end_row is not None and (not isinstance(data_end_row, int) or data_end_row < data_start_row):
            errors.append(f"{name} 数据结束行无效")
            continue
        try:
            path = resolve_current(suite, workbook_key)
        except Exception as exc:
            errors.append(f"{name} 当前工作簿无效：{exc}")
            continue
        book = load_workbook(path, read_only=True, data_only=False)
        try:
            if sheet_name not in book.sheetnames:
                errors.append(f"{name} Sheet不存在：{sheet_name}")
                continue
            sheet = book[sheet_name]
            headers = [cell.value for cell in sheet[header_row]]
            if key_header not in headers:
                errors.append(f"{name} 主键表头不存在：{key_header}")
                continue
            column = headers.index(key_header) + 1
            last_row = data_end_row if data_end_row is not None else (sheet.max_row or header_row)
            ids = [sheet.cell(row, column).value for row in range(data_start_row, last_row + 1)]
            populated = [str(value).strip() for value in ids if value not in (None, "")]
            duplicates = sorted(value for value, count in Counter(populated).items() if count > 1)
            if duplicates:
                errors.append(f"{name} 主键重复：{','.join(duplicates[:10])}")
            records += len(populated)
            details.append({
                "object": name, "workbook": workbook_key, "sheet": sheet_name,
                "key_header": key_header, "records": len(populated),
                "data_rows": [data_start_row, last_row],
                "source": path.relative_to(suite).as_posix(),
            })
        finally:
            book.close()
    return {"counts": {"objects": len(contracts), "records": records}, "errors": errors, "objects": details}
