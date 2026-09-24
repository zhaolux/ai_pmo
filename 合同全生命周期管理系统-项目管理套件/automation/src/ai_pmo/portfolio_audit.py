from __future__ import annotations

import re
from datetime import date, datetime
from pathlib import Path
from math import isclose
from openpyxl import load_workbook

CLOSED_STATES = ('已关闭', '已解决', '已取消')
SEVERE_DEFECT_LEVELS = ('Blocker', 'Critical', '致命', '严重')


def check_weekly_report(report: dict | None, dashboard_data_date=None,
                        today=None) -> dict:
    """周报发布口径：严重度嵌套、数据截止日一致性、与驾驶舱口径的滞后和周更节奏。

    report: parse_weekly_report 的解析结果；dashboard_data_date: AI PMO 数据接口
    的“数据日期”；today: 巡检当日（ISO 字符串或 date，测试可显式传入）。
    """
    if today is None:
        today = date.today()
    if isinstance(today, str):
        today = date.fromisoformat(today)
    if report is None:
        return {'errors': [], 'warnings': ['周报目录中未找到项目周报产物，发布口径无法核对']}
    errors = []
    warnings = []
    file_date = _iso_date(report.get('file_date'))
    data_date = _iso_date(report.get('data_date'))
    if data_date is None:
        warnings.append('最新周报缺少数据截止日，无法核对发布口径')
    else:
        if file_date is not None and data_date > file_date:
            errors.append(
                f'周报数据截止日 {data_date.isoformat()} 晚于发布日期 '
                f'{file_date.isoformat()}，口径不可能成立')
        dash = _iso_date(dashboard_data_date)
        if dash is not None and data_date < dash:
            warnings.append(
                f'周报数据截止日 {data_date.isoformat()} 早于驾驶舱数据日期 '
                f'{dash.isoformat()}，发布口径滞后于台账')
    counts = [report.get(key) for key in ('finding_count', 'major', 'high', 'medium')]
    if all(isinstance(c, int) for c in counts):
        finding_count, major, high, medium = counts
        if finding_count < major + high + medium:
            errors.append(
                f'周报严重度口径错误：发现总数 {finding_count} 小于重大+高+中合计 '
                f'{major + high + medium}')
    if file_date is not None and (today - file_date).days > 9:
        warnings.append(
            f'最新周报发布于 {file_date.isoformat()}，超过一周未更新，周更节奏中断')
    return {'errors': errors, 'warnings': warnings}


def parse_weekly_report(path: Path) -> dict | None:
    """从周报 xlsx 解析发布口径（文件名日期、数据截止日、健康度与严重度计数）。

    表结构变化时返回 None，由调用方按“无法核对”处理。
    """
    match = re.search(r'(\d{8})', Path(path).stem)
    file_date = _iso_date(match.group(1)) if match else None
    book = load_workbook(path, read_only=True, data_only=True)
    try:
        report: dict = {'file_date': file_date.isoformat() if file_date else None}
        if '使用说明' in book.sheetnames:
            for row in book['使用说明'].iter_rows(values_only=True):
                if row and row[0] == '数据截止日' and len(row) > 1 and row[1]:
                    parsed = _iso_date(row[1])
                    if parsed:
                        report['data_date'] = parsed.isoformat()
                    break
        if '周报摘要' not in book.sheetnames:
            return None
        sheet = book['周报摘要']
        rows = list(sheet.iter_rows(values_only=True))
        for index, row in enumerate(rows):
            cells = [str(v) if v is not None else '' for v in row]
            if '项目状态' in cells and '发现总数' in cells:
                cols = {name: cells.index(name) for name in
                        ('项目状态', '发现总数', '重大', '高', '中', '需人工确认')}
                data = rows[index + 1] if index + 1 < len(rows) else ()
                for name, col in cols.items():
                    key = {'项目状态': 'status', '发现总数': 'finding_count', '重大': 'major',
                           '高': 'high', '中': 'medium', '需人工确认': 'approval_count'}[name]
                    value = data[col] if col < len(data) else None
                    if key == 'status':
                        report[key] = str(value) if value not in (None, '') else None
                    else:
                        try:
                            report[key] = int(value)
                        except (TypeError, ValueError):
                            report[key] = None
                break
        report.setdefault('data_date', None)
        return report
    finally:
        book.close()


def _iso_date(value) -> date | None:
    if value in (None, ''):
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value).strip()
    for fmt in ('%Y-%m-%d', '%Y%m%d', '%Y/%m/%d'):
        try:
            return datetime.strptime(text[:10] if fmt == '%Y-%m-%d' else text, fmt).date()
        except ValueError:
            continue
    return None


def latest_weekly_report(weekly_dir: Path) -> Path | None:
    """按文件名日期（非修改时间）取最新周报 xlsx。"""
    candidates = []
    for path in Path(weekly_dir).glob('项目周报-*-V1.xlsx'):
        parsed = _iso_date(re.search(r'(\d{8})', path.stem).group(1)) \
            if re.search(r'(\d{8})', path.stem) else None
        if parsed:
            candidates.append((parsed, path))
    return max(candidates)[1] if candidates else None


def check_risk_followup(risks: list[dict], actions: dict[str, list[dict]]) -> dict:
    """未关闭且剩余等级为重大/高的风险，应对要素必须齐全。

    risks: [{'id', 'level'(剩余等级), 'status', 'strategy', 'deadline'}]；
    actions: {风险ID: [{'owner', 'plan_date'}]}（风险应对计划中的行动）。
    只提示缺失要素，不代替风险处置审批。
    """
    warnings = []
    for r in risks:
        if r['status'] in CLOSED_STATES or r['level'] not in ('重大', '高'):
            continue
        rid = r['id']
        if not r['strategy']:
            warnings.append(f'风险 {rid} 剩余等级{r["level"]}未关闭，但缺少应对策略')
        if not r['deadline']:
            warnings.append(f'风险 {rid} 剩余等级{r["level"]}未关闭，但缺少最迟应对日期')
        acts = actions.get(rid, [])
        if not acts:
            warnings.append(f'风险 {rid} 剩余等级{r["level"]}未关闭，风险应对计划中没有对应行动')
        for a in acts:
            if not a['owner']:
                warnings.append(f'风险 {rid} 的应对行动缺少行动责任人')
            if not a['plan_date']:
                warnings.append(f'风险 {rid} 的应对行动缺少计划完成日期')
    return {'errors': [], 'warnings': warnings}


def check_issue_detail(issues: list[dict]) -> dict:
    """问题台账状态明细：未关闭的重大/高问题须有应对措施与计划关闭日期。

    issues: [{'id', 'level', 'status', 'response', 'plan_close', 'actual_close'}]。
    """
    warnings = []
    for it in issues:
        closed = it['status'] in CLOSED_STATES
        if it['actual_close'] and not closed:
            warnings.append(
                f'问题 {it["id"]} 已填实际关闭日期，但状态为“{it["status"]}”，请确认是否已关闭')
        if closed or it['level'] not in ('重大', '高'):
            continue
        if not it['response']:
            warnings.append(f'问题 {it["id"]} 等级{it["level"]}未关闭，但缺少应对措施')
        if not it['plan_close']:
            warnings.append(f'问题 {it["id"]} 等级{it["level"]}未关闭，但缺少计划关闭日期')
    return {'errors': [], 'warnings': warnings}


def check_decision_detail(decisions: list[dict]) -> dict:
    """待决策事项：状态为待决策时，建议方案、需要日期、决策人必须齐全。

    decisions: [{'id', 'status', 'proposal', 'need_date', 'decision_maker'}]。
    """
    warnings = []
    for d in decisions:
        if d['status'] != '待决策':
            continue
        missing = []
        if not d['proposal']:
            missing.append('建议方案')
        if not d['need_date']:
            missing.append('需要日期')
        if not d['decision_maker']:
            missing.append('决策人')
        if missing:
            warnings.append(f'待决策事项 {d["id"]} 状态为待决策，但缺少{"、".join(missing)}')
    return {'errors': [], 'warnings': warnings}


def check_quality_evidence(coverage: list[dict] | None = None,
                           uat: list[dict] | None = None,
                           defects: list[dict] | None = None) -> dict:
    """质量验收证据一致性：判定为完整/通过时必须有证据，严重缺陷须有闭环计划。

    coverage: [{'id', 'exec_result', 'evidence', 'coverage'}]（需求覆盖矩阵）；
    uat: [{'id', 'exec_result', 'evidence'}]；
    defects: [{'id', 'severity', 'status', 'plan_done', 'verify_evidence'}]。
    """
    warnings = []
    for c in coverage or []:
        if c['coverage'] == '完整' and not c['evidence']:
            warnings.append(f'需求 {c["id"]} 覆盖状态为完整，但验收证据为空')
        if c['exec_result'] in ('通过', '已完成') and not c['evidence']:
            warnings.append(f'需求 {c["id"]} 执行结果为{c["exec_result"]}，但验收证据为空')
    for u in uat or []:
        if u['exec_result'] == '通过' and not u['evidence']:
            warnings.append(f'UAT {u["id"]} 执行结果为通过，但验收证据为空')
    for df in defects or []:
        if (df['severity'] in SEVERE_DEFECT_LEVELS
                and df['status'] not in CLOSED_STATES + ('已验证',)
                and not df['plan_done']):
            warnings.append(
                f'缺陷 {df["id"]} 严重级别{df["severity"]}未关闭，但缺少计划完成日期')
        if df['status'] in ('已关闭', '已解决') and not df['verify_evidence']:
            warnings.append(f'缺陷 {df["id"]} 已关闭，但验证证据为空')
    return {'errors': [], 'warnings': warnings}


def check_deliverable_consistency(rows: list[dict],
                                  interface_total: int | None = None) -> dict:
    """验收交付物台账口径与证据完整性。

    rows: [{'id', 'status', 'actual_done', 'review_result', 'location'}]；
    interface_total: 交付物工作簿 _数据接口登记的交付物总数。
    """
    warnings = []
    if interface_total is not None and interface_total != len(rows):
        warnings.append(
            f'交付物总数口径不一致：_数据接口登记 {interface_total}，台账 {len(rows)} 行')
    for d in rows:
        if d['status'] in ('已完成', '已验收') and not d['actual_done']:
            warnings.append(f'交付物 {d["id"]} 状态为{d["status"]}，但实际完成为空')
        if d['status'] == '已验收':
            missing = []
            if not d['review_result']:
                missing.append('评审结论')
            if not d['location']:
                missing.append('存放位置')
            if missing:
                warnings.append(f'交付物 {d["id"]} 已验收，但缺少{"、".join(missing)}')
    return {'errors': [], 'warnings': warnings}


def check_milestone_links(milestone_ids, task_milestones):
    milestones = set(milestone_ids)
    referenced = set(task_milestones)
    return {
        'errors': [f'总控任务引用不存在的里程碑 {item}' for item in sorted(referenced - milestones)],
        'warnings': [f'里程碑 {item} 没有总控任务' for item in sorted(milestones - referenced)],
    }


def compare_dashboard_snapshot(expected_sources, actual_sources, expected_metrics, actual_metrics):
    warnings = []
    for key, expected in expected_sources.items():
        actual = actual_sources.get(key)
        if actual != expected:
            warnings.append(f'驾驶舱来源快照过期：{key} 登记为{expected}，快照为{actual or "缺失"}')
    for key, expected in expected_metrics.items():
        if key not in actual_metrics or actual_metrics[key] in (None, ''):
            warnings.append(f'驾驶舱指标待刷新：{key} 源台账={expected}，快照缺失')
        elif (isinstance(expected, (int, float)) and isinstance(actual_metrics[key], (int, float))
              and isclose(float(expected), float(actual_metrics[key]), rel_tol=1e-9, abs_tol=0.005)):
            continue
        elif actual_metrics[key] != expected:
            warnings.append(f'驾驶舱指标不一致：{key} 源台账={expected}，快照={actual_metrics[key]}')
    return {'errors': [], 'warnings': warnings}


def dashboard_metrics_from_cells(values_by_row):
    rows = {
        '待决策事项': 14, '计划投入人天': 17, '计划人工成本': 18,
        '人日单价': 19, '里程碑总数': 36,
    }
    return {name: values_by_row.get(row) for name, row in rows.items()}


def check_migration_progress(interface: dict) -> dict:
    """Report migration source-inventory progress from the migration workbook interface."""
    clues = interface.get('来源线索数')
    inventoried = interface.get('已盘点来源系统数')
    gate_date = interface.get('G15 目标日期')
    if hasattr(gate_date, 'date'):
        gate_date = gate_date.date().isoformat()
    counts: dict = {}
    warnings: list[str] = []
    if clues in (None, ''):
        warnings.append('迁移工作簿缺少_数据接口盘点统计，无法核对来源系统盘点进度')
    else:
        counts = {'迁移来源线索': clues, '已盘点来源': inventoried}
        if inventoried != clues:
            warnings.append(
                f'迁移来源系统盘点未完成：已盘点 {inventoried or 0}/{clues}'
                + (f'，G15 目标日期 {gate_date}' if gate_date else '')
            )
    return {'counts': counts, 'warnings': warnings}


def check_milestone_owner_drift(plan_owners: dict[str, str],
                                dashboard_owners: dict[str, str]) -> dict:
    """里程碑负责人在计划工作簿与驾驶舱快照之间出现差异时提示。

    plan_owners / dashboard_owners: {里程碑名称: 负责人}。差异通常意味着
    里程碑计划刚被调整而驾驶舱尚未刷新，需要确认总控台账对应任务负责人已联动。
    """
    warnings = []
    for name, owner in plan_owners.items():
        snap = dashboard_owners.get(name)
        if snap and snap != owner:
            warnings.append(
                f'里程碑“{name}”负责人已变更：驾驶舱快照为{snap}，里程碑计划为{owner}；'
                '请确认总控台账对应任务负责人已联动，并刷新驾驶舱快照'
            )
    return {'errors': [], 'warnings': warnings}


def check_milestone_owner_alignment(milestones: list[dict],
                                    batch_owners: dict[str, list[str]]) -> dict:
    """里程碑负责人应直接负责其批次内至少一条任务。

    milestones: [{'id', 'name', 'owner', 'status'}]；batch_owners: {MS编号: [任务负责人…]}
    （任务负责人按业务→后端→前端→测试取第一个非空，空字符串表示该任务未落实负责人）。
    仅校验进行中/已完成且批次任务已全部落实负责人的里程碑，避免对未排期工作误报。
    """
    warnings = []
    for m in milestones:
        if m['status'] not in ('进行中', '已完成'):
            continue
        owners = batch_owners.get(m['id'], [])
        if not owners or any(not o for o in owners):
            continue
        if m['owner'] not in owners:
            warnings.append(
                f'里程碑 {m["id"]}“{m["name"]}”（负责人 {m["owner"]}）批次任务已全部落实负责人，'
                f'但无任何任务由 {m["owner"]} 负责，请核对里程碑负责人或任务负责人'
            )
    return {'errors': [], 'warnings': warnings}


def _dashboard_milestone_owners(sheet) -> dict[str, str]:
    """从驾驶舱项目驾驶舱表提取里程碑区块 {名称: 责任人}，表结构变化时返回空。"""
    try:
        rows = list(sheet.iter_rows(values_only=True))
    except Exception:
        return {}
    for i, row in enumerate(rows):
        cells = [str(v) if v is not None else '' for v in row]
        if '里程碑' in cells and '责任人' in cells:
            name_col, owner_col = cells.index('里程碑'), cells.index('责任人')
            owners: dict[str, str] = {}
            for data in rows[i + 1:]:
                name = data[name_col] if name_col < len(data) else None
                owner = data[owner_col] if owner_col < len(data) else None
                if name in (None, ''):
                    break
                owners[str(name)] = str(owner) if owner else ''
            return owners
    return {}


def _ids(sheet, prefix, start_row=1):
    return [str(row[0]) for row in sheet.iter_rows(min_row=start_row, values_only=True)
            if row and isinstance(row[0], str) and row[0].startswith(prefix)]


def _interface(sheet):
    return {str(row[0]): row[1] for row in sheet.iter_rows(values_only=True)
            if len(row) > 1 and row[0] not in (None, '')}


def audit_portfolio(plan_path: Path, risk_path: Path, change_path: Path,
                    quality_path: Path, deliverable_path: Path, cost_path: Path,
                    ai_pmo_path: Path, registered_sources: dict[str, str],
                    migration_path: Path | None = None,
                    weekly_dir: Path | None = None, today=None) -> dict:
    books = [load_workbook(path, read_only=True, data_only=True) for path in
             (plan_path, risk_path, change_path, quality_path, deliverable_path, cost_path, ai_pmo_path)]
    migration_book = load_workbook(migration_path, read_only=True, data_only=True) if migration_path else None
    if migration_book is not None:
        books.append(migration_book)
    plan, risk, change, quality, deliverable, cost, ai_pmo = books[:7]
    try:
        milestone_ids = _ids(plan['里程碑计划'], 'MS', 4)
        task_milestones = [str(row[5]) for row in plan['项目总控台账'].iter_rows(min_row=4, values_only=True)
                           if row and isinstance(row[0], str) and row[0].startswith('T') and row[5]]
        links = check_milestone_links(milestone_ids, task_milestones)
        ms_rows = []
        plan_owners: dict[str, str] = {}
        for row in plan['里程碑计划'].iter_rows(min_row=4, values_only=True):
            if row and isinstance(row[0], str) and row[0].startswith('MS'):
                ms_rows.append({'id': row[0], 'name': row[2], 'owner': row[4],
                                'status': row[10]})
                if row[2]:
                    plan_owners[str(row[2])] = str(row[4]) if row[4] else ''
        batch_owners: dict[str, list[str]] = {}
        for row in plan['项目总控台账'].iter_rows(min_row=4, values_only=True):
            if row and isinstance(row[0], str) and row[0].startswith('T') and row[5]:
                owner = next((str(v) for v in (row[11], row[15], row[19], row[23])
                              if v not in (None, '')), '')
                batch_owners.setdefault(str(row[5]), []).append(owner)
        owner_drift = check_milestone_owner_drift(
            plan_owners, _dashboard_milestone_owners(ai_pmo['项目驾驶舱']))
        owner_align = check_milestone_owner_alignment(ms_rows, batch_owners)
        interface = _interface(ai_pmo['_数据接口'])
        source_snapshot = {'plan': interface.get('里程碑来源'), 'risk': interface.get('风险来源')}
        risk_ids = _ids(risk['风险登记册'], 'RSK-', 9)
        decision_ids = _ids(change['待决策事项'], 'D-', 5)
        task_ids = _ids(plan['项目总控台账'], 'T', 4)
        cost_sheet = cost['人工成本预算']
        total_row = next((r for r in range(1, cost_sheet.max_row + 1)
                          if cost_sheet.cell(r, 1).value == '合计'), None)
        expected_metrics = {
            '总任务数': len(task_ids), '风险总数': len(risk_ids),
            '待决策事项': len(decision_ids), '里程碑总数': len(milestone_ids),
            '计划投入人天': cost_sheet.cell(total_row, 4).value if total_row else None,
            '计划人工成本': cost_sheet.cell(total_row, 6).value if total_row else None,
            '人日单价': cost_sheet['E3'].value,
        }
        actual_metrics = {name: interface.get(name) for name in ('总任务数', '风险总数')}
        actual_metrics.update(dashboard_metrics_from_cells({row: ai_pmo['_数据接口'].cell(row, 2).value
                                                             for row in (14, 17, 18, 19, 36)}))
        snapshot = compare_dashboard_snapshot(
            {key: value for key, value in registered_sources.items() if key in source_snapshot},
            source_snapshot, expected_metrics, actual_metrics,
        )
        counts = {
            '里程碑': len(milestone_ids), '总控任务': len(task_ids), '风险': len(risk_ids),
            '待决策': len(decision_ids), '缺陷': len(_ids(quality['缺陷台账'], 'DEF-', 5)),
            'UAT': len(_ids(quality['UAT验收'], 'UAT-', 5)),
            '交付物': len(_ids(deliverable['交付物台账'], 'D-', 5)),
        }
        migration = {'counts': {}, 'warnings': []}
        if migration_book is not None:
            migration = check_migration_progress(_interface(migration_book['_数据接口']))
            counts.update(migration['counts'])
        risk_rows = [{'id': row[0], 'level': row[33], 'status': row[28],
                      'strategy': row[24], 'deadline': row[27]}
                     for row in risk['风险登记册'].iter_rows(min_row=9, values_only=True)
                     if row and isinstance(row[0], str) and row[0].startswith('RSK-')]
        risk_actions: dict[str, list[dict]] = {}
        for row in risk['风险应对计划'].iter_rows(min_row=5, values_only=True):
            if row and isinstance(row[0], str) and row[0].startswith('RA-') and row[1]:
                risk_actions.setdefault(str(row[1]), []).append(
                    {'owner': row[6], 'plan_date': row[9]})
        risk_followup = check_risk_followup(risk_rows, risk_actions)
        issue_detail = check_issue_detail([
            {'id': row[0], 'level': row[3], 'status': row[10],
             'response': row[6], 'plan_close': row[7], 'actual_close': row[8]}
            for row in change['问题台账'].iter_rows(min_row=4, values_only=True)
            if row and row[0] not in (None, '') and str(row[0]) != '编号'])
        decision_detail = check_decision_detail([
            {'id': row[0], 'status': row[8], 'proposal': row[5],
             'need_date': row[7], 'decision_maker': row[6]}
            for row in change['待决策事项'].iter_rows(min_row=5, values_only=True)
            if row and isinstance(row[0], str) and row[0].startswith('D-')])
        quality_evidence = check_quality_evidence(
            coverage=[{'id': row[0], 'exec_result': row[6], 'evidence': row[9],
                       'coverage': row[11]}
                      for row in quality['需求覆盖矩阵'].iter_rows(min_row=5, values_only=True)
                      if row and isinstance(row[0], str) and row[0].startswith('REQ-')],
            uat=[{'id': row[0], 'exec_result': row[9], 'evidence': row[11]}
                 for row in quality['UAT验收'].iter_rows(min_row=5, values_only=True)
                 if row and isinstance(row[0], str) and row[0].startswith('UAT-')],
            defects=[{'id': row[0], 'severity': row[6], 'status': row[13],
                      'plan_done': row[11], 'verify_evidence': row[16]}
                     for row in quality['缺陷台账'].iter_rows(min_row=5, values_only=True)
                     if row and isinstance(row[0], str) and row[0].startswith('DEF-')])
        deliverable_rows = [{'id': row[0], 'status': row[6], 'actual_done': row[5],
                             'review_result': row[8], 'location': row[10]}
                            for row in deliverable['交付物台账'].iter_rows(min_row=5, values_only=True)
                            if row and isinstance(row[0], str) and row[0].startswith('D-')]
        deliverable_interface = _interface(deliverable['_数据接口'])
        deliverable_check = check_deliverable_consistency(
            deliverable_rows, interface_total=deliverable_interface.get('交付物总数'))
        detail_warnings = (risk_followup['warnings'] + issue_detail['warnings']
                           + decision_detail['warnings'] + quality_evidence['warnings']
                           + deliverable_check['warnings'])
        weekly = {'errors': [], 'warnings': []}
        if weekly_dir is not None:
            latest = latest_weekly_report(weekly_dir)
            parsed = parse_weekly_report(latest) if latest else None
            weekly = check_weekly_report(
                parsed, dashboard_data_date=interface.get('数据日期'), today=today)
        return {'counts': counts, 'errors': links['errors'] + snapshot['errors'] + owner_drift['errors'] + owner_align['errors'] + weekly['errors'],
                'warnings': links['warnings'] + snapshot['warnings'] + migration['warnings']
                + owner_drift['warnings'] + owner_align['warnings'] + detail_warnings
                + weekly['warnings']}
    finally:
        for book in books: book.close()
