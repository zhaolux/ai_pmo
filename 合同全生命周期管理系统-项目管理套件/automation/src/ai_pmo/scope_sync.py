"""03 需求与范围管理 ↔ 04 系统集成管理 接口快照双向同步。

列归属（见 00_使用说明与配置/数据入口与联动审计）：
- 范围字段（接口名称、方向、范围决策、优先级、责任人、规格定义等）以 03 为权威，同步 03→04；
- 执行字段（规格状态、测试状态、备注）以 04 为权威，同步 04→03。

新增记录以 03 为准追加到 04；04 出现 03 没有的 ID 视为结构错误，不删除、不写入。
"""
from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

INTERFACE_TABLE = ("接口清单", "接口清单")
SPEC_TABLE = ("P0接口规格台账", "P0接口规格")

INTERFACE_SCOPE_COLUMNS = (
    "清单ID", "集成系统", "接口名称", "接口描述", "接口方向",
    "范围决策", "优先级", "接口责任人",
)
INTERFACE_EXECUTION_COLUMNS = ("规格状态", "测试状态", "备注")

SPEC_SCOPE_COLUMNS = (
    "规格ID", "清单ID", "集成系统", "接口名称", "业务描述", "优先级", "接口方向",
    "源系统", "目标系统", "交互模式", "调用频率", "关键字段", "鉴权方式", "幂等键",
    "超时/重试", "异常补偿", "对账规则", "数据敏感级别", "接口责任人", "计划评审日期",
)
SPEC_EXECUTION_COLUMNS = ("规格状态", "备注")

HEADER_ROW = 4
DATA_START_ROW = 5


@dataclass(frozen=True)
class TableSyncPlan:
    source_sheet: str
    copy_sheet: str
    scope_columns: tuple[str, ...]
    execution_columns: tuple[str, ...]
    source_rows: int = 0
    copy_rows: int = 0
    scope_updates: tuple[str, ...] = ()      # 03→04 的范围字段差异（记录ID列表）
    execution_updates: tuple[str, ...] = ()  # 04→03 的执行字段差异（记录ID列表）
    appends: tuple[str, ...] = ()            # 03 有而 04 无、需追加的记录ID
    errors: tuple[str, ...] = ()

    @property
    def has_changes(self) -> bool:
        return bool(self.scope_updates or self.execution_updates or self.appends)


@dataclass
class SyncReport:
    scope_path: Path
    integration_path: Path
    applied: bool
    plans: list[TableSyncPlan] = field(default_factory=list)

    def to_dict(self) -> dict:
        return {
            "scope": self.scope_path.name,
            "integration": self.integration_path.name,
            "applied": self.applied,
            "tables": [
                {
                    "source_sheet": p.source_sheet,
                    "copy_sheet": p.copy_sheet,
                    "source_rows": p.source_rows,
                    "copy_rows": p.copy_rows,
                    "scope_updates": len(p.scope_updates),
                    "execution_updates": len(p.execution_updates),
                    "appends": len(p.appends),
                    "errors": list(p.errors),
                }
                for p in self.plans
            ],
        }


def _columns(sheet) -> list[str]:
    for row in sheet.iter_rows(min_row=HEADER_ROW, max_row=HEADER_ROW, values_only=True):
        return [str(value) if value is not None else "" for value in row]
    return []


def _records(sheet, columns: list[str]) -> dict[str, dict[str, object]]:
    records: dict[str, dict[str, object]] = {}
    for row in sheet.iter_rows(min_row=DATA_START_ROW, values_only=True):
        key = str(row[0]).strip() if row and row[0] is not None else ""
        if not key:
            continue
        records[key] = {columns[i]: row[i] if i < len(row) else None for i in range(len(columns))}
    return records


def _plan_table(source, copy, source_sheet: str, copy_sheet: str,
                scope_columns: tuple[str, ...], execution_columns: tuple[str, ...]) -> TableSyncPlan:
    source_cols = _columns(source)
    copy_cols = _columns(copy)
    for column in scope_columns + execution_columns:
        if column not in source_cols or column not in copy_cols:
            return TableSyncPlan(
                source_sheet, copy_sheet, scope_columns, execution_columns,
                errors=(f"列缺失：{column} 不在两表表头中",),
            )
    source_records = _records(source, source_cols)
    copy_records = _records(copy, copy_cols)
    errors: list[str] = []
    for key in sorted(copy_records.keys() - source_records.keys()):
        errors.append(f"{copy_sheet}:{key} 存在于04系统集成但不在03需求范围，未删除未写入，请人工核实")
    scope_updates = tuple(
        key for key in sorted(source_records.keys() & copy_records.keys())
        if any(source_records[key].get(column) != copy_records[key].get(column) for column in scope_columns)
    )
    execution_updates = tuple(
        key for key in sorted(source_records.keys() & copy_records.keys())
        if any(source_records[key].get(column) != copy_records[key].get(column) for column in execution_columns)
    )
    return TableSyncPlan(
        source_sheet, copy_sheet, scope_columns, execution_columns,
        source_rows=len(source_records), copy_rows=len(copy_records),
        scope_updates=scope_updates, execution_updates=execution_updates,
        appends=tuple(sorted(source_records.keys() - copy_records.keys())),
        errors=tuple(errors),
    )


def plan_sync(scope_path: Path, integration_path: Path) -> SyncReport:
    scope = load_workbook(scope_path, data_only=False, read_only=True)
    integration = load_workbook(integration_path, data_only=False, read_only=True)
    try:
        plans = [
            _plan_table(
                scope[INTERFACE_TABLE[0]], integration[INTERFACE_TABLE[1]],
                INTERFACE_TABLE[0], INTERFACE_TABLE[1],
                INTERFACE_SCOPE_COLUMNS, INTERFACE_EXECUTION_COLUMNS,
            ),
            _plan_table(
                scope[SPEC_TABLE[0]], integration[SPEC_TABLE[1]],
                SPEC_TABLE[0], SPEC_TABLE[1],
                SPEC_SCOPE_COLUMNS, SPEC_EXECUTION_COLUMNS,
            ),
        ]
    finally:
        scope.close()
        integration.close()
    return SyncReport(scope_path, integration_path, applied=False, plans=plans)


def _write_values(sheet, records: dict[str, dict[str, object]], columns: tuple[str, ...],
                  values_by_key: dict[str, dict[str, object]]) -> int:
    """按记录ID把 values_by_key 中对应列写回 sheet，返回写入单元格数。"""
    header = _columns(sheet)
    column_positions = {name: header.index(name) for name in columns}
    written = 0
    for row_number, row in enumerate(sheet.iter_rows(min_row=DATA_START_ROW), start=DATA_START_ROW):
        key = str(row[0].value).strip() if row[0].value is not None else ""
        if key not in values_by_key:
            continue
        for name, position in column_positions.items():
            value = values_by_key[key].get(name)
            cell = row[position]
            if cell.value != value:
                cell.value = value
                written += 1
    return written


def _append_records(sheet, records: dict[str, dict[str, object]],
                    keys: tuple[str, ...], columns: list[str]) -> int:
    """把 records 中 keys 指定的记录追加到 sheet 末尾，沿用末行样式。"""
    if not keys:
        return 0
    last_row = max(sheet.max_row, DATA_START_ROW - 1)
    template = [sheet.cell(row=last_row, column=i + 1) for i in range(len(columns))]
    appended = 0
    for offset, key in enumerate(keys, start=1):
        target_row = last_row + offset
        for position, name in enumerate(columns):
            cell = sheet.cell(row=target_row, column=position + 1)
            cell._style = template[position]._style
            cell.value = records[key].get(name)
        appended += 1
    return appended


def apply_sync(scope_path: Path, integration_path: Path) -> SyncReport:
    report = plan_sync(scope_path, integration_path)
    if any(plan.errors for plan in report.plans):
        return report
    if not any(plan.has_changes for plan in report.plans):
        return report

    scope = load_workbook(scope_path, data_only=False)
    integration = load_workbook(integration_path, data_only=False)
    try:
        for plan in report.plans:
            source = scope[plan.source_sheet]
            copy = integration[plan.copy_sheet]
            source_records = _records(source, _columns(source))
            copy_records = _records(copy, _columns(copy))

            _write_values(
                copy, copy_records, plan.scope_columns,
                {key: source_records[key] for key in plan.scope_updates + plan.appends},
            )
            _write_values(
                source, source_records, plan.execution_columns,
                {key: copy_records[key] for key in plan.execution_updates},
            )
            _append_records(copy, source_records, plan.appends, _columns(copy))
        scope.save(scope_path)
        integration.save(integration_path)
    finally:
        scope.close()
        integration.close()

    report.applied = True
    verify = plan_sync(scope_path, integration_path)
    report.plans = [
        TableSyncPlan(
            p.source_sheet, p.copy_sheet, p.scope_columns, p.execution_columns,
            source_rows=p.source_rows, copy_rows=p.copy_rows,
            errors=tuple(p.errors) + (
                ("写入后复核仍存在差异",) if (p.scope_updates or p.execution_updates or p.appends) else ()
            ),
        )
        for p in verify.plans
    ]
    return report


def column_letter(name: str, columns: tuple[str, ...]) -> str:
    return get_column_letter(columns.index(name) + 1)
