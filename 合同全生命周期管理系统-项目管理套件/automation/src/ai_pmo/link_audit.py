from __future__ import annotations

import re
from collections.abc import Iterable, Sequence
from pathlib import Path

from openpyxl import load_workbook


Row = tuple[int, tuple[object, ...]]


def _index(rows: Iterable[Row], label: str, side: str) -> tuple[dict[str, Row], list[str]]:
    indexed: dict[str, Row] = {}
    findings: list[str] = []
    for row_number, values in rows:
        key = str(values[0] or "").strip()
        if not key:
            continue
        if key in indexed:
            findings.append(f"{label}:{side}重复ID {key} 行{indexed[key][0]}/{row_number}")
        else:
            indexed[key] = (row_number, values)
    return indexed, findings


def compare_snapshots(
    source: Iterable[Row], copy: Iterable[Row], label: str, columns: Sequence[str]
) -> list[str]:
    left, findings = _index(source, label, "源表")
    right, duplicates = _index(copy, label, "副本")
    findings.extend(duplicates)
    for key in sorted(left.keys() - right.keys()):
        findings.append(f"{label}:{key} 副本缺失 源行{left[key][0]}")
    for key in sorted(right.keys() - left.keys()):
        findings.append(f"{label}:{key} 源表缺失 副本行{right[key][0]}")
    for key in sorted(left.keys() & right.keys()):
        source_row, source_values = left[key]
        copy_row, copy_values = right[key]
        for position, name in enumerate(columns):
            if position == 0:
                continue
            a = source_values[position] if position < len(source_values) else None
            b = copy_values[position] if position < len(copy_values) else None
            if a != b:
                findings.append(
                    f"{label}:{key} {name}不一致 源行{source_row}/副本行{copy_row}"
                )
    return findings


def audit_interface_relations(
    interface_ids: set[str], specs: Iterable[Row],
    implementations: Iterable[Row], task_ids: set[str]
) -> list[str]:
    specs_by_id, findings = _index(specs, "P0规格", "表内")
    impl_by_id, impl_duplicates = _index(implementations, "实施任务", "表内")
    findings.extend(impl_duplicates)
    for spec_id, (row_number, values) in specs_by_id.items():
        interface_id = str(values[1] or "").strip()
        if interface_id not in interface_ids:
            findings.append(f"P0规格:{spec_id} 引用不存在的清单ID {interface_id} 行{row_number}")
    for impl_id, (row_number, values) in impl_by_id.items():
        spec_id = str(values[1] or "").strip()
        interface_id = str(values[2] or "").strip()
        task_id = str(values[3] or "").strip()
        if spec_id not in specs_by_id:
            findings.append(f"实施任务:{impl_id} 引用不存在的规格ID {spec_id} 行{row_number}")
        elif str(specs_by_id[spec_id][1][1] or "").strip() != interface_id:
            expected = str(specs_by_id[spec_id][1][1] or "").strip()
            findings.append(
                f"实施任务:{impl_id} 规格{spec_id}所属接口{expected}与任务接口{interface_id}不一致 行{row_number}"
            )
        if interface_id not in interface_ids:
            findings.append(f"实施任务:{impl_id} 引用不存在的清单ID {interface_id} 行{row_number}")
        if task_id not in task_ids:
            findings.append(f"实施任务:{impl_id} 引用不存在的总控任务 {task_id} 行{row_number}")
    return findings


def resource_task_ids(texts: Iterable[str]) -> set[str]:
    return {task_id for text in texts for task_id in re.findall(r"\bT\d{3}\b", text or "")}


def _records(sheet, header_row: int, columns: Sequence[str], prefix: str) -> list[Row]:
    headers = [cell.value for cell in sheet[header_row]]
    positions = [headers.index(column) for column in columns]
    result: list[Row] = []
    for row_number, values in enumerate(sheet.iter_rows(min_row=header_row + 1, values_only=True), header_row + 1):
        selected = tuple(values[pos] if pos < len(values) else None for pos in positions)
        if isinstance(selected[0], str) and selected[0].startswith(prefix):
            result.append((row_number, selected))
    return result


def audit_suite_links(scope_path: Path, integration_path: Path, plan_path: Path) -> dict[str, object]:
    """Read source workbooks only; report snapshot drift and unresolved resource coverage."""
    scope = load_workbook(scope_path, read_only=True, data_only=True)
    integration = load_workbook(integration_path, read_only=True, data_only=True)
    plan = load_workbook(plan_path, read_only=True, data_only=True)
    try:
        interfaces_columns = tuple(cell.value for cell in scope["接口清单"][4] if cell.value)
        specs_columns = tuple(cell.value for cell in scope["P0接口规格台账"][4] if cell.value)
        source_interfaces = _records(scope["接口清单"], 4, interfaces_columns, "IF-")
        copy_interfaces = _records(integration["接口清单"], 4, interfaces_columns, "IF-")
        source_specs = _records(scope["P0接口规格台账"], 4, specs_columns, "SPEC-")
        copy_specs = _records(integration["P0接口规格"], 4, specs_columns, "SPEC-")
        implementation = _records(
            integration["接口实施计划"], 4,
            ("实施任务ID", "规格ID", "清单ID", "关联总控任务"), "IMP-",
        )
        task_rows = _records(plan["项目总控台账"], 3, ("任务ID",), "T")
        task_ids = {str(values[0]) for _, values in task_rows}
        assigned = resource_task_ids(
            str(row[8] or "") for row in plan["资源计划"].iter_rows(min_row=39, values_only=True)
        )
        errors = compare_snapshots(source_interfaces, copy_interfaces, "接口清单", interfaces_columns)
        errors += compare_snapshots(source_specs, copy_specs, "P0规格", specs_columns)
        errors += audit_interface_relations(
            {str(values[0]) for _, values in source_interfaces}, source_specs, implementation, task_ids,
        )
        warnings = [f"资源计划:总控任务 {task_id} 未出现于人员周任务清单" for task_id in sorted(task_ids - assigned)]
        errors += [f"资源计划:任务清单引用不存在的总控任务 {task_id}" for task_id in sorted(assigned - task_ids)]
        return {
            "counts": {
                "接口清单": len(source_interfaces), "P0规格": len(source_specs),
                "实施任务": len(implementation), "总控任务": len(task_ids),
                "资源已覆盖任务": len(task_ids & assigned),
            },
            "errors": errors,
            "warnings": warnings,
        }
    finally:
        scope.close()
        integration.close()
        plan.close()
