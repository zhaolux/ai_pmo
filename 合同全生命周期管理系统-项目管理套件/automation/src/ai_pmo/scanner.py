from __future__ import annotations

import hashlib
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import load_workbook

from .config import Paths, load_json
from .database import connect


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def scan(paths: Paths) -> dict[str, int]:
    mapping = load_json("sheet_mapping.json")
    reverse_mapping = {
        sheet: code for code, sheet_names in mapping.items() for sheet in sheet_names
    }
    file_count = 0
    sheet_count = 0
    scanned_at = datetime.now(timezone.utc).isoformat()
    candidates = sorted(
        path for path in paths.inputs.rglob("*")
        if path.is_file() and path.suffix.lower() in {".xlsx", ".xlsm", ".docx", ".csv"}
    )
    with connect(paths.database) as connection:
        for path in candidates:
            stat = path.stat()
            relative = path.relative_to(paths.suite).as_posix()
            connection.execute(
                """INSERT INTO source_file(relative_path, file_type, size_bytes, modified_at, sha256, scanned_at)
                   VALUES(?, ?, ?, ?, ?, ?)
                   ON CONFLICT(relative_path) DO UPDATE SET
                     file_type=excluded.file_type,
                     size_bytes=excluded.size_bytes,
                     modified_at=excluded.modified_at,
                     sha256=excluded.sha256,
                     scanned_at=excluded.scanned_at""",
                (relative, path.suffix.lower(), stat.st_size,
                 datetime.fromtimestamp(stat.st_mtime, timezone.utc).isoformat(),
                 sha256(path), scanned_at),
            )
            source_id = connection.execute(
                "SELECT source_id FROM source_file WHERE relative_path=?", (relative,)
            ).fetchone()[0]
            file_count += 1
            if path.suffix.lower() in {".xlsx", ".xlsm"}:
                workbook = load_workbook(path, read_only=True, data_only=False)
                connection.execute("DELETE FROM source_sheet WHERE source_id=?", (source_id,))
                for order, sheet in enumerate(workbook.worksheets, start=1):
                    connection.execute(
                        """INSERT INTO source_sheet
                           (source_id, sheet_name, sheet_order, max_row, max_column, module_code)
                           VALUES(?, ?, ?, ?, ?, ?)""",
                        (source_id, sheet.title, order, sheet.max_row, sheet.max_column,
                         reverse_mapping.get(sheet.title)),
                    )
                    sheet_count += 1
                workbook.close()
    return {"files": file_count, "sheets": sheet_count}

