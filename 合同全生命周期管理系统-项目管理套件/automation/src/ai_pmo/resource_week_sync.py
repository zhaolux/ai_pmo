"""资源计划周任务清单与项目总控台账联动刷新（默认只读，显式执行才写工作簿）。

联动口径：人员周计划的"任务清单"是人员→任务→周的人工分配权威（保持不变），
本模块按总控台账当前任务名称/状态刷新清单文本、任务数与周工时，
并重算 master 汇总（任务数/未完成任务/最早开始/最晚结束/计划投入工时）。
清单引用了总控不存在的任务ID时视为结构错误，拒绝写入。
"""
from __future__ import annotations

import re
from dataclasses import dataclass, field
from pathlib import Path

from openpyxl import load_workbook

CONTROL_SHEET = "项目总控台账"
RESOURCE_SHEET = "资源计划"
MASTER_START = 5
MASTER_END = 34
WEEK_START_ROW = 39
TASK_ID_RE = re.compile(r"T\d{3}")
HOURS_PER_WEEK = 40


@dataclass
class ResourceSyncReport:
    rows_scanned: int = 0
    rows_refreshed: int = 0
    persons: list[str] = field(default_factory=list)
    errors: list[str] = field(default_factory=list)
    applied: bool = False

    def to_dict(self) -> dict:
        return {
            "rows_scanned": self.rows_scanned,
            "rows_refreshed": self.rows_refreshed,
            "persons": self.persons,
            "errors": self.errors,
            "applied": self.applied,
        }


def read_control_tasks(plan_path: Path) -> dict[str, dict]:
    """任务ID → 名称/总体状态（数据仅来自总控台账）。"""
    workbook = load_workbook(plan_path, data_only=True)
    try:
        control = workbook[CONTROL_SHEET]
        tasks: dict[str, dict] = {}
        for row in control.iter_rows(min_row=4, values_only=True):
            if row[0]:
                tasks[str(row[0])] = {
                    "name": str(row[4] or "").strip(),
                    "status": str(row[6] or "").strip(),
                }
        return tasks
    finally:
        workbook.close()


def plan_resource_sync(plan_path: Path) -> ResourceSyncReport:
    tasks = read_control_tasks(plan_path)
    workbook = load_workbook(plan_path, data_only=False)
    try:
        resource = workbook[RESOURCE_SHEET]
        report = ResourceSyncReport()
        missing: set[str] = set()
        for row in resource.iter_rows(min_row=WEEK_START_ROW, values_only=True):
            report.rows_scanned += 1
            task_ids = TASK_ID_RE.findall(str(row[8] or ""))
            if not task_ids:
                continue
            report.rows_refreshed += 1
            for task_id in task_ids:
                if task_id not in tasks:
                    missing.add(task_id)
        persons = [
            str(resource.cell(row=r, column=1).value)
            for r in range(MASTER_START, MASTER_END + 1)
            if resource.cell(row=r, column=1).value
        ]
        report.persons = persons
        if missing:
            report.errors.append(
                "资源计划任务清单引用了总控台账不存在的任务：" + "、".join(sorted(missing))
            )
        return report
    finally:
        workbook.close()


def apply_resource_sync(plan_path: Path) -> ResourceSyncReport:
    report = plan_resource_sync(plan_path)
    if report.errors:
        return report
    tasks = read_control_tasks(plan_path)

    workbook = load_workbook(plan_path)
    try:
        resource = workbook[RESOURCE_SHEET]
        # 第一遍：刷新周行任务清单/任务数/周工时，并按人聚合
        person_weeks: dict[str, list[int]] = {}
        current_person = ""
        for row_number in range(WEEK_START_ROW, resource.max_row + 1):
            name = resource.cell(row=row_number, column=1).value
            if name:
                current_person = str(name)
            task_ids = TASK_ID_RE.findall(str(resource.cell(row=row_number, column=9).value or ""))
            if not task_ids:
                continue
            text = "；".join(f"{task_id} {tasks[task_id]['name']}" for task_id in task_ids)
            resource.cell(row=row_number, column=9).value = text
            resource.cell(row=row_number, column=4).value = "有"
            resource.cell(row=row_number, column=5).value = len(task_ids)
            resource.cell(row=row_number, column=6).value = HOURS_PER_WEEK
            person_weeks.setdefault(current_person, []).append(row_number)

        # 第二遍：重算 master 汇总
        for row_number in range(MASTER_START, MASTER_END + 1):
            person = resource.cell(row=row_number, column=1).value
            if not person:
                continue
            person = str(person)
            rows = person_weeks.get(person, [])
            task_ids: list[str] = []
            for week_row in rows:
                task_ids += TASK_ID_RE.findall(str(resource.cell(row=week_row, column=9).value or ""))
            distinct = list(dict.fromkeys(task_ids))
            open_tasks = [t for t in distinct if tasks[t]["status"] != "已完成"]
            resource.cell(row=row_number, column=3).value = len(distinct)
            resource.cell(row=row_number, column=4).value = len(open_tasks)
            if rows:
                starts = [resource.cell(row=r, column=2).value for r in rows]
                finishes = [resource.cell(row=r, column=3).value for r in rows]
                resource.cell(row=row_number, column=5).value = min(starts)
                resource.cell(row=row_number, column=6).value = max(finishes)
            resource.cell(row=row_number, column=8).value = HOURS_PER_WEEK * len(rows)

        workbook.save(plan_path)
        report.applied = True
        return report
    finally:
        workbook.close()
