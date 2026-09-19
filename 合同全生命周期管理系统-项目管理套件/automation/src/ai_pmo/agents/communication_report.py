from __future__ import annotations

from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook

from ..agent_models import AgentResult, Evidence, Finding


def _date(value) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    return None


def analyze_communication_report(communication_file: Path, as_of: date) -> AgentResult:
    workbook = load_workbook(communication_file, data_only=True, read_only=True)
    findings: list[Finding] = []
    try:
        actions = workbook["行动项台账"]
        for row_number, row in enumerate(actions.iter_rows(min_row=5, values_only=True), start=5):
            if not row[0]:
                continue
            status = str(row[9] or "")
            if status in {"已完成", "已关闭", "已取消"}:
                continue
            due = _date(row[7])
            if not due or due >= as_of:
                continue
            action_id = str(row[0])
            findings.append(Finding(
                finding_id=f"COM-ACTION-{action_id}", agent="communication_report",
                object_id=action_id, severity="高", title="行动项已逾期",
                detail=f"{row[3] or action_id}截止日为{due.isoformat()}，当前状态为{status or '空'}。",
                recommendation="确认完成承诺、阻塞原因、协同人和可验证的关闭证据。",
                owner=str(row[4] or "未指定"), requires_approval=True,
                evidence=(Evidence(communication_file.name, actions.title, action_id, row_number),),
            ))

        decisions = workbook["决策日志"]
        for row_number, row in enumerate(decisions.iter_rows(min_row=5, values_only=True), start=5):
            if not row[0]:
                continue
            status = str(row[9] or "")
            if status in {"已决策", "已关闭", "已取消"}:
                continue
            due = _date(row[4])
            if not due:
                continue
            days = (due - as_of).days
            if days > 14:
                continue
            decision_id = str(row[0])
            overdue = days < 0
            findings.append(Finding(
                finding_id=f"COM-DECISION-{decision_id}", agent="communication_report",
                object_id=decision_id, severity="高" if overdue else "中",
                title="决策事项已逾期" if overdue else "决策事项14天内到期",
                detail=f"{row[1] or decision_id}决策截止日为{due.isoformat()}，当前状态为{status or '空'}。",
                recommendation="将事实、备选方案、影响、建议和所需授权统一提交决策人。",
                owner=str(row[10] or "未指定"), requires_approval=True,
                evidence=(Evidence(communication_file.name, decisions.title, decision_id, row_number),),
            ))

        escalations = workbook["升级事项"]
        for row_number, row in enumerate(escalations.iter_rows(min_row=5, values_only=True), start=5):
            if not row[0]:
                continue
            status = str(row[8] or "")
            if status in {"已关闭", "已解决", "已取消"}:
                continue
            escalation_id = str(row[0])
            due = _date(row[7])
            overdue = bool(due and due < as_of)
            findings.append(Finding(
                finding_id=f"COM-ESCALATION-{escalation_id}", agent="communication_report",
                object_id=escalation_id, severity="重大" if overdue else "高",
                title="升级事项待处理", detail=str(row[3] or escalation_id),
                recommendation="明确所需决策或支持、接收人、截止日和处理结论。",
                owner=str(row[10] or row[9] or "未指定"), requires_approval=True,
                evidence=(Evidence(communication_file.name, escalations.title, escalation_id, row_number),),
            ))
    finally:
        workbook.close()
    return AgentResult("communication_report", f"识别{len(findings)}项沟通、决策或升级关注事项", tuple(findings))
