from __future__ import annotations

from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook

from ..agent_models import AgentResult, Evidence, Finding


def _date(value) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    return None


def _closed(status: object) -> bool:
    return str(status or "").strip() in {"已关闭", "已验证", "已解决", "不修复", "已取消"}


def analyze_quality_acceptance(quality_file: Path, deliverable_file: Path, as_of: date) -> AgentResult:
    workbook = load_workbook(quality_file, data_only=True, read_only=True)
    deliverable_book = load_workbook(deliverable_file, data_only=True, read_only=True)
    findings: list[Finding] = []
    try:
        overview = workbook["质量总览"]
        for row_number, row in enumerate(overview.iter_rows(min_row=10, values_only=True), start=10):
            if not row[0]:
                continue
            status = str(row[1] or "")
            due = _date(row[3])
            if status in {"已通过", "通过", "已关闭"} or not due:
                continue
            days = (due - as_of).days
            if days > 14:
                continue
            overdue = days < 0
            gate = str(row[0])
            findings.append(Finding(
                finding_id=f"QUA-GATE-{row_number}", agent="quality_acceptance",
                object_id=f"GATE-{gate}", severity="高" if overdue else "中",
                title="质量门禁已逾期" if overdue else "质量门禁14天内到期",
                detail=f"{gate}计划日期为{due.isoformat()}，当前状态为{status or '空'}。",
                recommendation="核对通过标准、责任人和证据位置，完成评审或提交升级决策。",
                owner=str(row[2] or "未指定"), requires_approval=overdue,
                evidence=(Evidence(quality_file.name, overview.title, gate, row_number),),
            ))

        defects = workbook["缺陷台账"]
        for row_number, row in enumerate(defects.iter_rows(min_row=5, values_only=True), start=5):
            if not row[0] or _closed(row[13] if len(row) > 13 else None):
                continue
            level = str(row[6] or "")
            if level not in {"Blocker", "Critical", "阻断", "严重", "重大", "高"}:
                continue
            defect_id = str(row[0])
            major = level in {"Blocker", "Critical", "阻断", "重大"}
            findings.append(Finding(
                finding_id=f"QUA-DEFECT-{defect_id}", agent="quality_acceptance",
                object_id=defect_id, severity="重大" if major else "高",
                title="未关闭严重缺陷", detail=f"{row[5] or defect_id}；当前状态为{row[13] or '空'}。",
                recommendation="明确根因、修复版本、完成日、回归范围和验证证据。",
                owner=str(row[8] or "未指定"), requires_approval=major,
                evidence=(Evidence(quality_file.name, defects.title, defect_id, row_number),),
            ))

        deliverables = deliverable_book["交付物台账"]
        for row_number, row in enumerate(deliverables.iter_rows(min_row=5, values_only=True), start=5):
            if not row[0]:
                continue
            due = _date(row[4])
            status = str(row[6] or "")
            if not due or status in {"已完成", "已验收", "已批准"}:
                continue
            days = (due - as_of).days
            if days > 14:
                continue
            deliverable_id = str(row[0])
            overdue = days < 0
            findings.append(Finding(
                finding_id=f"QUA-DELIVERABLE-{deliverable_id}", agent="quality_acceptance",
                object_id=deliverable_id, severity="高" if overdue else "中",
                title="质量交付物已逾期" if overdue else "质量交付物14天内到期",
                detail=f"{row[1] or deliverable_id}计划日期为{due.isoformat()}，当前状态为{status or '空'}。",
                recommendation="确认版本、评审人、评审结论和存放位置。",
                owner=str(row[3] or "未指定"), requires_approval=overdue,
                evidence=(Evidence(deliverable_file.name, deliverables.title, deliverable_id, row_number),),
            ))

        uat = workbook["UAT验收"]
        for row_number, row in enumerate(uat.iter_rows(min_row=5, values_only=True), start=5):
            if not row[0]:
                continue
            due = _date(row[7])
            result = str(row[9] or "")
            evidence_value = row[11] if len(row) > 11 else None
            confirmed = str(row[12] or "") if len(row) > 12 else ""
            if not due or due > as_of or (result in {"通过", "已通过"} and evidence_value and confirmed in {"已确认", "通过"}):
                continue
            uat_id = str(row[0])
            findings.append(Finding(
                finding_id=f"QUA-UAT-{uat_id}", agent="quality_acceptance",
                object_id=uat_id, severity="高", title="UAT到期但验收证据不完整",
                detail=f"{row[1] or uat_id}计划执行日为{due.isoformat()}，执行结果为{result or '空'}。",
                recommendation="补齐执行人、结果、遗留风险、验收证据和业务负责人确认。",
                owner="产品经理", requires_approval=True,
                evidence=(Evidence(quality_file.name, uat.title, uat_id, row_number),),
            ))
    finally:
        deliverable_book.close()
        workbook.close()
    return AgentResult("quality_acceptance", f"识别{len(findings)}项质量与验收关注事项", tuple(findings))
