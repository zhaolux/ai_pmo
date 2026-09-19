from __future__ import annotations

from datetime import date
from pathlib import Path

from openpyxl import load_workbook

from ..agent_models import AgentResult, Evidence, Finding


def _rows(sheet, start: int):
    return ((number, row) for number, row in enumerate(
        sheet.iter_rows(min_row=start, values_only=True), start=start
    ) if row[0])


def analyze_risk_issue(risk_file: Path, change_file: Path, as_of: date) -> AgentResult:
    findings: list[Finding] = []
    risk_book = load_workbook(risk_file, data_only=True, read_only=True)
    try:
        sheet = risk_book["风险登记册"]
        for row_number, row in _rows(sheet, 9):
            level = str(row[33] or "")
            if level not in {"重大", "高"}:
                continue
            risk_id, name = str(row[0]), str(row[1] or "")
            findings.append(Finding(
                finding_id=f"RSK-{risk_id}", agent="risk_issue", object_id=risk_id,
                severity=level, title=f"{level}剩余风险", detail=name,
                recommendation="复核触发条件、应对行动、剩余风险和升级结论。",
                owner=str(row[22] or "未指定"), requires_approval=level == "重大",
                evidence=(Evidence(risk_file.name, sheet.title, risk_id, row_number),),
            ))
    finally:
        risk_book.close()

    change_book = load_workbook(change_file, data_only=True, read_only=True)
    try:
        issue_sheet = change_book["问题台账"]
        for row_number, row in _rows(issue_sheet, 4):
            issue_id = str(row[0])
            findings.append(Finding(
                finding_id=f"ISS-{issue_id}", agent="risk_issue", object_id=issue_id,
                severity=str(row[3] or "中"), title="未关闭问题",
                detail=str(row[2] or row[1] or ""), recommendation="确认根因、行动、期限和关闭证据。",
                owner=str(row[5] or "未指定"), requires_approval=False,
                evidence=(Evidence(change_file.name, issue_sheet.title, issue_id, row_number),),
            ))
        change_sheet = change_book["变更台账"]
        for row_number, row in _rows(change_sheet, 4):
            change_id = str(row[0])
            if str(row[6] or "") in {"已批准", "已关闭", "已拒绝"}:
                continue
            findings.append(Finding(
                finding_id=f"CHG-{change_id}", agent="risk_issue", object_id=change_id,
                severity=str(row[5] or "中"), title="待审批变更",
                detail=str(row[3] or ""), recommendation="完成范围、进度、成本和风险影响评估后提交审批。",
                owner=str(row[2] or "未指定"), requires_approval=True,
                evidence=(Evidence(change_file.name, change_sheet.title, change_id, row_number),),
            ))
        decision_sheet = change_book["待决策事项"]
        for row_number, row in _rows(decision_sheet, 5):
            decision_id = str(row[0])
            if str(row[8] or "") in {"已决策", "已关闭"}:
                continue
            findings.append(Finding(
                finding_id=f"DEC-{decision_id}", agent="risk_issue", object_id=decision_id,
                severity="高", title="待决策事项", detail=str(row[1] or ""),
                recommendation="明确决策人、最迟日期、备选方案和延迟后果。",
                owner=str(row[4] or "未指定"), requires_approval=True,
                evidence=(Evidence(change_file.name, decision_sheet.title, decision_id, row_number),),
            ))
    finally:
        change_book.close()
    return AgentResult("risk_issue", f"识别{len(findings)}项风险、问题、变更或决策事项", tuple(findings))
