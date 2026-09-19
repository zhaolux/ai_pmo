from __future__ import annotations

from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook

from ..agent_models import AgentResult, Evidence, Finding


def _date(value) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    return None


def analyze_cost_contract(cost_file: Path, as_of: date) -> AgentResult:
    workbook = load_workbook(cost_file, data_only=True, read_only=True)
    findings: list[Finding] = []
    try:
        baseline = workbook["预算基线"]
        missing_budget: list[tuple[int, str, str, str]] = []
        for row_number, row in enumerate(
            baseline.iter_rows(min_row=5, values_only=True), start=5
        ):
            if not row[0]:
                continue
            status = str(row[11] or "")
            if status != "已量化":
                missing_budget.append((row_number, str(row[0]), str(row[2] or ""), str(row[10] or "未指定")))
        if missing_budget:
            row_number, budget_id, _, owner = missing_budget[0]
            names = "、".join(item[2] for item in missing_budget[:5])
            suffix = "等" if len(missing_budget) > 5 else ""
            findings.append(Finding(
                finding_id="COST-BUDGET-INCOMPLETE", agent="cost_contract",
                object_id="BUDGET-COMPLETENESS", severity="高",
                title="项目总预算尚未完整量化",
                detail=f"共{len(missing_budget)}个预算项仍待估算：{names}{suffix}。",
                recommendation="完成报价、估算依据和预备费评估，经授权人批准后形成总预算基线。",
                owner=owner, requires_approval=True,
                evidence=(Evidence(cost_file.name, baseline.title, budget_id, row_number),),
            ))

        packages = workbook["采购包计划"]
        for row_number, row in enumerate(
            packages.iter_rows(min_row=5, values_only=True), start=5
        ):
            if not row[0]:
                continue
            package_id = str(row[0])
            due = _date(row[5])
            amount = row[7]
            status = str(row[8] or "")
            if amount not in (None, "") or status not in {"待估算", "", "未开始"} or not due:
                continue
            days = (due - as_of).days
            if days > 14:
                continue
            overdue = days < 0
            findings.append(Finding(
                finding_id=f"COST-{package_id}-ESTIMATE", agent="cost_contract",
                object_id=package_id, severity="高" if overdue else "中",
                title="采购包估算已逾期" if overdue else "采购包估算14天内到期",
                detail=f"{row[1] or package_id}估算完成日为{due.isoformat()}，计划金额仍为空。",
                recommendation="确认计价边界、可比报价、估算依据和预算纳入时点。",
                owner=str(row[4] or "未指定"), requires_approval=overdue,
                evidence=(Evidence(cost_file.name, packages.title, package_id, row_number),),
            ))

        if "付款里程碑" in workbook.sheetnames:
            payments = workbook["付款里程碑"]
            for row_number, row in enumerate(
                payments.iter_rows(min_row=5, values_only=True), start=5
            ):
                if not row[0]:
                    continue
                payment_id = str(row[0])
                actual_amount = row[13] if len(row) > 13 else None
                evidence = row[8] if len(row) > 8 else None
                acceptance = str(row[9] or "") if len(row) > 9 else ""
                if actual_amount not in (None, "", 0) and (not evidence or acceptance not in {"已验收", "通过"}):
                    findings.append(Finding(
                        finding_id=f"COST-{payment_id}-EVIDENCE", agent="cost_contract",
                        object_id=payment_id, severity="高", title="付款缺少验收证据",
                        detail=f"已登记实际付款{actual_amount}元，但验收证据或验收状态不完整。",
                        recommendation="补齐交付物、验收结论、发票和付款审批证据。",
                        owner="项目经理", requires_approval=True,
                        evidence=(Evidence(cost_file.name, payments.title, payment_id, row_number),),
                    ))
    finally:
        workbook.close()
    return AgentResult("cost_contract", f"识别{len(findings)}项成本与合同关注事项", tuple(findings))
