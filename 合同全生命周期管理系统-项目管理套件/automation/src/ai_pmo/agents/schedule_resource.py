from __future__ import annotations

from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook

from ..agent_models import AgentResult, Evidence, Finding


def _date(value) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    return None


def analyze_schedule_resource(plan_file: Path, as_of: date) -> AgentResult:
    workbook = load_workbook(plan_file, data_only=True, read_only=True)
    findings: list[Finding] = []
    try:
        sheet = workbook["项目总控台账"]
        for row_number, row in enumerate(sheet.iter_rows(min_row=4, values_only=True), start=4):
            if not row[0]:
                continue
            task_id = str(row[0])
            task_name = str(row[4] or "")
            status = str(row[6] or "")
            owner = str(row[7] or "未指定")
            finish = _date(row[9])
            if finish and finish < as_of and status != "已完成":
                findings.append(Finding(
                    finding_id=f"SCH-{task_id}-OVERDUE", agent="schedule_resource",
                    object_id=task_id, severity="高", title="任务已延期",
                    detail=f"{task_name}计划完成日为{finish.isoformat()}，当前状态为{status or '空'}。",
                    recommendation="确认实际状态、纠偏计划、承诺日期和完成证据。",
                    owner=owner, requires_approval=True,
                    evidence=(Evidence(plan_file.name, sheet.title, task_id, row_number),),
                ))
            elif finish and status != "已完成" and 0 <= (finish - as_of).days <= 14:
                findings.append(Finding(
                    finding_id=f"SCH-{task_id}-DUE14", agent="schedule_resource",
                    object_id=task_id, severity="中", title="任务14天内到期",
                    detail=f"{task_name}将在{finish.isoformat()}到期，当前状态为{status or '空'}。",
                    recommendation="确认交付物、验收人和按期完成证据。",
                    owner=owner, requires_approval=False,
                    evidence=(Evidence(plan_file.name, sheet.title, task_id, row_number),),
                ))
        if "资源计划" in workbook.sheetnames:
            resource = workbook["资源计划"]
            person = ""
            for row_number, row in enumerate(resource.iter_rows(min_row=39, values_only=True), start=39):
                if row[0]:
                    person = str(row[0]).strip()
                week_start = _date(row[1]) if len(row) > 1 else None
                week_finish = _date(row[2]) if len(row) > 2 else None
                load_rate = row[6] if len(row) > 6 else None
                judgment = str(row[7] or "") if len(row) > 7 else ""
                if not person or not week_start or not week_finish or not (week_start <= as_of <= week_finish):
                    continue
                overloaded = judgment == "超负荷" or isinstance(load_rate, (int, float)) and load_rate > 1
                if overloaded:
                    week_key = week_start.strftime("%Y%m%d")
                    task_list = str(row[8] or "未列明") if len(row) > 8 else "未列明"
                    findings.append(Finding(
                        finding_id=f"RES-{person}-{week_key}-OVERLOAD", agent="schedule_resource",
                        object_id=person, severity="高", title="人员当前周超负荷",
                        detail=f"{person}在{week_start.isoformat()}当周负荷率为{load_rate}，任务为{task_list}。",
                        recommendation="调整任务分配或明确加班审批，确保单人周计划不超过40小时。",
                        owner=person, requires_approval=True,
                        evidence=(Evidence(plan_file.name, resource.title, person, row_number),),
                    ))
    finally:
        workbook.close()
    return AgentResult("schedule_resource", f"识别{len(findings)}项进度关注事项", tuple(findings))
