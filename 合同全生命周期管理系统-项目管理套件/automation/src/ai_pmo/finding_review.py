from __future__ import annotations

from collections import Counter
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import load_workbook

from .database import connect


ALLOWED_STATUSES = {"待确认", "已接受", "需调整", "已关闭"}


def _read_review_workbook(path: Path) -> tuple[str, list[dict]]:
    book = load_workbook(path, read_only=True, data_only=True)
    try:
        if "使用说明" not in book.sheetnames or "发现清单" not in book.sheetnames:
            raise ValueError("复核工作簿缺少使用说明或发现清单")
        run_id = ""
        for row in book["使用说明"].iter_rows(values_only=True):
            if row and str(row[0] or "").strip() == "运行ID":
                run_id = str(row[1] or "").strip()
                break
        if not run_id:
            raise ValueError("使用说明中缺少运行ID")
        sheet = book["发现清单"]
        rows: list[dict] = []
        for row in sheet.iter_rows(min_row=5, values_only=True):
            finding_id = str(row[0] or "").strip()
            status = str(row[10] or "").strip() if len(row) > 10 else ""
            conclusion = str(row[11] or "").strip() if len(row) > 11 else ""
            if not finding_id:
                continue
            rows.append({"finding_id": finding_id, "status": status or "待确认", "conclusion": conclusion})
        return run_id, rows
    finally:
        book.close()


def import_reviews(database: Path, workbook: Path, reviewer: str) -> dict:
    reviewer = reviewer.strip()
    if not reviewer:
        raise ValueError("复核人不能为空")
    run_id, rows = _read_review_workbook(workbook)
    duplicates = [item for item, count in Counter(row["finding_id"] for row in rows).items() if count > 1]
    if duplicates:
        raise ValueError(f"发现ID重复：{','.join(duplicates)}")
    invalid_statuses = sorted({row["status"] for row in rows if row["status"] not in ALLOWED_STATUSES})
    if invalid_statuses:
        raise ValueError(f"确认状态无效：{','.join(invalid_statuses)}")
    selected = [row for row in rows if row["status"] != "待确认"]
    if any(not row["conclusion"] for row in selected):
        raise ValueError("非待确认记录必须填写人工结论")
    with connect(database) as connection:
        run = connection.execute("SELECT 1 FROM agent_run WHERE run_id=?", (run_id,)).fetchone()
        if not run:
            raise ValueError(f"运行ID不存在：{run_id}")
        valid_ids = {row[0] for row in connection.execute(
            "SELECT finding_id FROM agent_finding WHERE run_id=?", (run_id,)
        )}
        unknown = sorted(row["finding_id"] for row in rows if row["finding_id"] not in valid_ids)
        if unknown:
            raise ValueError(f"发现ID不存在：{','.join(unknown)}")
        now = datetime.now(timezone.utc).isoformat()
        for row in selected:
            connection.execute(
                """INSERT INTO agent_finding_review
                   (run_id, finding_id, status, conclusion, reviewer, reviewed_at, source_path)
                   VALUES(?,?,?,?,?,?,?)
                   ON CONFLICT(run_id, finding_id) DO UPDATE SET
                     status=excluded.status, conclusion=excluded.conclusion,
                     reviewer=excluded.reviewer, reviewed_at=excluded.reviewed_at,
                     source_path=excluded.source_path""",
                (run_id, row["finding_id"], row["status"], row["conclusion"], reviewer, now, str(workbook)),
            )
    return {"run_id": run_id, "imported": len(selected), "reviewer": reviewer}


def review_summary(database: Path, run_id: str) -> dict:
    counts = {status: 0 for status in ALLOWED_STATUSES}
    with connect(database) as connection:
        rows = connection.execute(
            """SELECT COALESCE(r.status, '待确认') AS status, COUNT(*) AS count
               FROM agent_finding f
               LEFT JOIN agent_finding_review r
                 ON r.run_id=f.run_id AND r.finding_id=f.finding_id
               WHERE f.run_id=? GROUP BY COALESCE(r.status, '待确认')""",
            (run_id,),
        )
        for row in rows:
            counts[row["status"]] = row["count"]
    return counts
