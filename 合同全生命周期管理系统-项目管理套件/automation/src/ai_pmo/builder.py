from __future__ import annotations

from copy import copy, deepcopy
from datetime import date
from pathlib import Path

from openpyxl import Workbook, load_workbook

from .config import Paths, load_json
from .output_guard import ensure_outputs_available


MODULE_NAMES = {
    "00": "使用说明与配置", "01": "项目启动与治理", "02": "计划与进度管理",
    "03": "需求与范围管理", "04": "系统集成管理", "05": "成本与合同管理",
    "06": "风险问题与变更", "07": "质量测试与验收", "08": "沟通会议与报告",
    "09": "上线移交与运营", "10": "复盘与经验沉淀", "11": "AI PMO中心"
}


def _copy_sheet(source, target) -> None:
    for row in source.iter_rows():
        for source_cell in row:
            target_cell = target[source_cell.coordinate]
            target_cell.value = source_cell.value
            if source_cell.has_style:
                # Workbook style indexes are local; copy the public style parts.
                target_cell.font = copy(source_cell.font)
                target_cell.fill = copy(source_cell.fill)
                target_cell.border = copy(source_cell.border)
                target_cell.alignment = copy(source_cell.alignment)
                target_cell.protection = copy(source_cell.protection)
            if source_cell.number_format:
                target_cell.number_format = source_cell.number_format
            if source_cell.hyperlink:
                target_cell._hyperlink = copy(source_cell.hyperlink)
            if source_cell.comment:
                target_cell.comment = copy(source_cell.comment)
    for key, dimension in source.column_dimensions.items():
        target.column_dimensions[key].width = dimension.width
        target.column_dimensions[key].hidden = dimension.hidden
    for key, dimension in source.row_dimensions.items():
        target.row_dimensions[key].height = dimension.height
        target.row_dimensions[key].hidden = dimension.hidden
    for merged in source.merged_cells.ranges:
        target.merge_cells(str(merged))
    target.freeze_panes = source.freeze_panes
    target.auto_filter.ref = source.auto_filter.ref
    target.sheet_view.showGridLines = source.sheet_view.showGridLines
    target.sheet_properties = copy(source.sheet_properties)
    target.page_margins = copy(source.page_margins)
    target.page_setup = copy(source.page_setup)
    target.print_options = copy(source.print_options)
    target.sheet_format = copy(source.sheet_format)
    for validation in source.data_validations.dataValidation:
        target.add_data_validation(deepcopy(validation))
    for image in source._images:
        target.add_image(deepcopy(image))
    for chart in source._charts:
        target.add_chart(deepcopy(chart))


def _add_interface_sheet(workbook: Workbook, project: dict, module_code: str) -> None:
    sheet = workbook.create_sheet("_数据接口")
    sheet.append(["字段", "值"])
    sheet.append(["Project_ID", project["project_id"]])
    sheet.append(["项目名称", project["project_name"]])
    sheet.append(["模块编码", module_code])
    sheet.append(["生成日期", date.today().isoformat()])
    sheet.append(["中央数据库", "../90_项目数据中心/ai_pmo.db"])
    sheet.sheet_state = "hidden"


def build_output_paths(
    paths: Paths, directories: dict | None = None,
) -> dict[str, Path]:
    directories = directories or load_json("directories.json")
    date_text = date.today().strftime("%Y%m%d")
    return {
        code: paths.suite / directories[code] / f"{module_name}-{date_text}-V1.xlsx"
        for code, module_name in MODULE_NAMES.items()
    }


def build(paths: Paths) -> list[Path]:
    project = load_json("project.json")
    directories = load_json("directories.json")
    mapping = load_json("sheet_mapping.json")
    template = paths.inputs / "参考资料" / project["template_filename"]
    if not template.exists():
        raise FileNotFoundError(f"缺少V2参考模板：{template}")
    targets = build_output_paths(paths, directories)
    ensure_outputs_available(targets.values())
    source = load_workbook(template, data_only=False)
    outputs: list[Path] = []
    for code, module_name in MODULE_NAMES.items():
        workbook = Workbook()
        workbook.remove(workbook.active)
        copied = 0
        for sheet_name in mapping.get(code, []):
            if sheet_name not in source.sheetnames:
                continue
            target = workbook.create_sheet(sheet_name)
            _copy_sheet(source[sheet_name], target)
            copied += 1
        if copied == 0:
            sheet = workbook.create_sheet("模块说明")
            sheet.append([project["project_name"]])
            sheet.append(["模块", module_name])
            sheet.append(["状态", "待补充标准模板"])
            sheet.column_dimensions["A"].width = 28
            sheet.column_dimensions["B"].width = 36
        _add_interface_sheet(workbook, project, code)
        target_dir = paths.suite / directories[code]
        target_dir.mkdir(parents=True, exist_ok=True)
        target_path = targets[code]
        workbook.save(target_path)
        outputs.append(target_path)
    source.close()
    return outputs
